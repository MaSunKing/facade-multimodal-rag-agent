"""15 additional simple English CPU cases; freeze gold before pipeline execution.

Official documents are reused, but cells/clauses differ from previous questions.
Conflict controls are explicitly fictional native TXT pairs, not official labels.
No GPU weights, answer generation, or production changes.
"""
from __future__ import annotations
import hashlib
import argparse
import json
import os
from pathlib import Path
import re
import shutil
import sys
import time

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT/'scripts'))
from run_fresh_english_native_eval import norm, support

FOLDER = ROOT/'evaluation/basic_capabilities_extra_20260917'
OUTPUT = ROOT/'outputs/public_eval/basic_capabilities_extra_20260917'
ORIGINALS = ROOT/'evaluation/fresh_english_native_20260917/assets'


def save(path, obj):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def cases():
    rows = []
    def add(kind, question, **gold):
        rows.append(dict(id=f'extra_{kind}_{sum(r["category"]==kind for r in rows)+1:02}',
                         category=kind, question=question, **gold))
    for cell, value, label, unit_cell, unit, question in [
        ('D12',774873,'Nuclear','A6','Thousand Megawatthours','What was nuclear net generation across all sectors in 2023?'),
        ('C15',451904,'Wind','A6','Thousand Megawatthours','What was wind net generation across all sectors in 2024?'),
        ('C24',83918,'Estimated Small Scale Solar Photovoltaic','A6','Thousand Megawatthours','What was estimated small scale solar photovoltaic generation across all sectors in 2024?'),
        ('C28',372931,'Coal','A28','1000 tons','How much coal was consumed for electricity generation across all sectors in 2024?'),
        ('B52',3975382,'All Sectors','B46','million kWh','What were total sales of electricity to ultimate customers across all sectors in 2024?'),
    ]:
        add('numeric', question, file='eia_industry_summary.xlsx', anchors=[str(value)],
            row_label=label, native_locator=dict(sheet='epa_01_01',cell=cell,value=value),
            unit=unit,unit_cell=unit_cell, provenance='project_authored_native_cell_verified')
    for file, paragraph, anchors, question in [
        ('ofgem_householder_application.docx',27,['criterion 1','cannot be used with','criterion 3'],'Can criterion 1 be used together with criterion 3 in the householder application?'),
        ('ofgem_householder_application.docx',65,['Route 4','ECO4 Flex only'],'Which scheme is Route 4 applicable to, and is its scope restricted?'),
        ('ofgem_householder_application.docx',67,['To use Route 4','new method','approved by the Department for Energy Security and Net Zero'],'What approval is needed for a new household identification method before Route 4 can be used?'),
        ('ofgem_medical_referral.docx',14,['Route 2','other than their low-income','supplemented by a second eligibility criteria by the local authority'],'Under Route 2, what additional eligibility step supplements a medical referral for a reason other than low income?'),
        ('ofgem_medical_referral.docx',34,['Printed on headed paper','practice, clinic or hospital making the referral'],'What paper and referring organisation information are required for a printed medical referral?'),
    ]:
        add('condition', question, file=file, anchors=anchors,
            native_locator=dict(paragraph=paragraph),provenance='project_authored_native_clause_verified')
    for entity, metric, a, b, unit in [
        ('MockTile-Q','thickness','17','19','mm'),
        ('MockStrip-R','width','80','85','mm'),
        ('MockBeam-S','length','2.4','2.7','m'),
        ('MockPanel-T','weight','4.2','4.6','kg'),
        ('MockSheet-U','area','1.2','1.4','m2'),
    ]:
        add('conflict',f'Do the two sources agree on the {metric} of {entity}? Report both measurements.',
            entity=entity, metric=metric, values=[a,b],unit=unit,
            files=[f'{entity}_source_a.txt',f'{entity}_source_b.txt'],
            provenance='project_authored_fictional_same_scope_disagreement_control')
    return rows


def prepare_and_audit(rows):
    from docx import Document
    from openpyxl import load_workbook
    assets=FOLDER/'assets'; assets.mkdir(parents=True,exist_ok=True)
    names=sorted({r['file'] for r in rows if 'file' in r})
    for name in names:
        target=assets/name
        if target.exists() and sha(target)!=sha(ORIGINALS/name): raise ValueError('Source changed')
        if not target.exists(): shutil.copy2(ORIGINALS/name,target)
    audit=[]
    for r in rows:
        if r['category']=='conflict':
            for name,value in zip(r['files'],r['values']):
                text=('Synthetic evaluation data, not a real product specification.\n'
                      'Measurement scope: same laboratory sample and same batch.\n'
                      f'{r["entity"]} {r["metric"]}: {value} {r["unit"]}.\n')
                path=assets/name
                if path.exists() and path.read_text(encoding='utf-8')!=text: raise ValueError('Frozen fixture differs')
                if not path.exists(): path.write_text(text,encoding='utf-8')
            audit.append(dict(id=r['id'],valid=True,source_kind='synthetic',facts=list(zip(r['files'],r['values']))))
        elif r['category']=='numeric':
            loc=r['native_locator']; book=load_workbook(assets/r['file'],read_only=True,data_only=True)
            sheet=book[loc['sheet']]; value=sheet[loc['cell']].value
            unit_text=str(sheet[r['unit_cell']].value)
            row=int(re.search(r'\d+',loc['cell']).group()); row_label=str(sheet[f'A{row}'].value)
            valid=value==loc['value'] and norm(r['unit']) in norm(unit_text) and norm(r['row_label']) in norm(row_label)
            audit.append(dict(id=r['id'],valid=valid,cell=loc['cell'],value=value,row_label=row_label,unit_source=unit_text))
            book.close()
        else:
            text=Document(assets/r['file']).paragraphs[r['native_locator']['paragraph']].text
            audit.append(dict(id=r['id'],valid=all(norm(a) in norm(text) for a in r['anchors']),source_excerpt=text))
    if not all(a['valid'] for a in audit): raise ValueError('Independent source audit failed')
    frozen=FOLDER/'questions.json'
    if frozen.exists() and json.loads(frozen.read_text(encoding='utf-8'))!=rows: raise ValueError('Questions changed')
    save(frozen,rows)
    save(FOLDER/'source_manifest.json',[dict(file=p.name,sha256=sha(p),source_kind='synthetic' if p.suffix=='.txt' else 'official_document_reused_new_fact') for p in sorted(assets.iterdir())])
    save(OUTPUT/'independent_gold_audit.json',audit)


def numeric_flags(items,r):
    matching=[i for i in items if support(i,r)]
    unit_match=False
    # Require the native unit cell (not an unrelated unit mention) in the same
    # workbook/sheet context as the retained answer cell. Do not inject gold.
    for i in items:
        if i.get('document_name')!=r['file'] or f'sheet={r["native_locator"]["sheet"]}' not in i.get('text',''): continue
        text=i.get('text','')
        if re.search(r'(?<![A-Z0-9])'+r['unit_cell']+r'(?:\[[^\]]*\])?=',text) and norm(r['unit']) in norm(text): unit_match=True
    return dict(target_row_value_retained=bool(matching),native_unit_retained=bool(matching) and unit_match,
                row_value_unit_retained=bool(matching) and unit_match)


def condition_flags(items,r):
    matches=[i for i in items if support(i,r)]
    return dict(condition_relation_retained=bool(matches),
                condition_protection_tagged=any('condition' in i.get('protected_relation_types',[]) for i in matches))


def conflict_support(items,r):
    return all(any(i.get('document_name')==file and norm(f'{r["entity"]} {r["metric"]}: {value} {r["unit"]}') in norm(i.get('text',''))
                   for i in items) for file,value in zip(r['files'],r['values']))


def main():
    global OUTPUT
    parser=argparse.ArgumentParser()
    parser.add_argument('--prompt-budget',type=int,default=5000)
    args=parser.parse_args()
    OUTPUT=OUTPUT/f'budget_{args.prompt_budget}'
    rows=cases(); prepare_and_audit(rows)
    for key in ['CUSTOMER_DOCUMENT_OCR_ENABLED','CUSTOMER_ATTACHMENT_SEMANTIC_RERANK','RAG_HYBRID_ENABLED']: os.environ[key]='0'
    from backend.documents.customer_sessions import add_files,bind_session_owner,delete_session,retrieve
    from backend.sales.context_engine import optimise_evidence_context,validate_packed_evidence
    from backend.app import compact_grounded_payload_for_generation
    from transformers import AutoTokenizer
    tokenizer=AutoTokenizer.from_pretrained(os.getenv('PUBLIC_EVAL_TOKENIZER_PATH',str(ROOT/'models/Qwen3-VL-8B-Instruct')),local_files_only=True)
    owner='extra-basic-cpu'; sessions={}; results=[]
    try:
        for r in rows:
            start=time.perf_counter(); out=dict(id=r['id'],category=r['category'],question=r['question'])
            try:
                names=tuple(r.get('files') or [r['file']])
                if names not in sessions:
                    sessions[names]=add_files([(n,(FOLDER/'assets'/n).read_bytes()) for n in names],owner_id=owner)['session_id']
                with bind_session_owner(owner): recalled=retrieve(sessions[names],r['question'],max_chunks=24,max_text_tokens=5000)
                ranked,context=optimise_evidence_context(r['question'],recalled['evidence'])
                payload,packing=compact_grounded_payload_for_generation(dict(customer_question=r['question'],evidence=ranked),tokenizer,
                                    max_prompt_tokens=args.prompt_budget,system_prompt='Answer only from evidence.')
                packed=json.loads(payload)['evidence']; back={i['evidence_id']:i for i in ranked}
                scored=[dict(back.get(i['evidence_id'],{}),**i) for i in packed]
                if r['category']=='numeric':
                    out.update(numeric_flags(scored,r),before_packing=numeric_flags(ranked,r))
                    out['passed']=out['row_value_unit_retained']
                elif r['category']=='condition':
                    out.update(condition_flags(scored,r),before_packing=condition_flags(ranked,r))
                    out['passed']=out['condition_relation_retained']
                else:
                    # Detection must bind the intended entity, metric, BOTH source
                    # files and BOTH complete values, not an arbitrary conflict.
                    groups=[]
                    for group in context.get('conflict_groups',[]):
                        members=[back[x] for x in group.get('evidence_ids',[]) if x in back]
                        if group.get('entity')==r['entity'].casefold() and group.get('metric')==r['metric'] and conflict_support(members,r): groups.append(group)
                    out.update(conflict_detected=bool(groups),both_sources_retained=conflict_support(scored,r),
                               before_packing=dict(both_sources_retained=conflict_support(ranked,r)),matching_conflict_groups=groups)
                    out['passed']=out['conflict_detected'] and out['both_sources_retained']
                out.update(ranked_evidence=ranked,packed_evidence=packed,context_audit=context,packing_audit=packing,
                           integrity=validate_packed_evidence(ranked,packed),retrieval_audit=recalled['input_snapshot'])
            except Exception as exc:
                out.update(passed=False,error=dict(type=type(exc).__name__,message=str(exc)))
            out['elapsed_seconds']=round(time.perf_counter()-start,3); results.append(out)
            print(json.dumps({k:v for k,v in out.items() if k not in ['ranked_evidence','packed_evidence','context_audit','packing_audit','retrieval_audit','integrity']},ensure_ascii=True),flush=True)
    finally:
        for sid in sessions.values(): delete_session(sid,owner_id=owner)
    fields={'numeric':['target_row_value_retained','native_unit_retained','row_value_unit_retained'],
            'condition':['condition_relation_retained','condition_protection_tagged'],
            'conflict':['conflict_detected','both_sources_retained']}
    metrics={kind:dict(cases=5,passed=sum(r.get('passed',False) for r in results if r['category']==kind),
                      **{field:dict(hits=sum(r.get(field,False) for r in results if r['category']==kind),total=5) for field in keys}) for kind,keys in fields.items()}
    report=dict(track='native_file_CPU_retrieval_context_exact_tokenizer_packing',model_generation_tested=False,
                retrieval_token_budget=5000,packing_prompt_budget=args.prompt_budget,questions_sha256=sha(FOLDER/'questions.json'),
                official_documents_reused=3,new_synthetic_documents=10,new_questions=15,metrics=metrics,details=results)
    save(OUTPUT/'results.json',report)
    print(json.dumps(metrics),flush=True)


if __name__=='__main__': main()
