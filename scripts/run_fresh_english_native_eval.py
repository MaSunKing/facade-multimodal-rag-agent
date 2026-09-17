"""New official native files, English source-grounded questions; CPU track only.

Gold answers/locators below are authored from independent pypdf/docx/openpyxl
inspection, not retrieval predictions. No paraphrase augmentation or dropped fails.
"""
from __future__ import annotations
import hashlib
import json
import os
from pathlib import Path
import re
import sys
import time

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
FOLDER = ROOT / 'evaluation/fresh_english_native_20260917'
OUTPUT = ROOT / 'outputs/public_eval/fresh_english_native_20260917'


def norm(text):
    return ''.join(str(text).split()).casefold()


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def cases():
    rows = []
    def add(file, question, answer, anchors, **locator):
        rows.append(dict(id=f'fresh_en_{len(rows)+1:03}', file=file,
                         question=question, expected_answer=answer,
                         anchors=anchors, native_locator=locator,
                         label_provenance='project_authored_independently_source_verified_not_official_QA'))
    add('ofgem_data_dictionary.pdf', 'In what file format must the declaration notification be submitted?', '.csv', ['file is submitted', '.csv'], page=2)
    add('ofgem_data_dictionary.pdf', 'What does the first section of the LA unique reference number represent?', 'The 9-digit ONS code assigned to the local authority.', ['9-digit', 'Statistics (ONS)', 'assigned'], page=3)
    add('ofgem_data_dictionary.pdf', 'How many digits must an ECO4 Flex Route 4 application number contain?', '5 digits', ['unique 5-digit', 'application'], page=6)
    add('ofgem_third_party_facts.pdf', 'Does an ECO4 or GBIS Flex third-party referral guarantee that measures will be installed?', 'No.', ['does not guarantee', 'measures will be installed'], page=2)
    add('ofgem_third_party_facts.pdf', 'What signed document must Citizens Advice produce when referring a household struggling to pay energy bills?', 'A signed referral letter to the local authority confirming the evidence of difficulty paying gas or electricity bills.', ['signed referral letter', 'struggling to pay'], page=2)
    add('ofgem_third_party_facts.pdf', 'For Flex Route 2 Proxy 3, is an NHS referral letter mandatory?', 'No; alternative evidence such as hospital diagnosis letters or prescriptions can be used.', ['NHS referral', 'not mandatory', 'alternative evidence'], page=4)
    add('gmfrs_information_box.pdf', 'What colours should the lobby sector and fire sector folders be?', 'Lobby: yellow; fire: red.', ['yellow in colour', 'red in colour'], page=3)
    add('gmfrs_information_box.pdf', 'How many copies of the Building Emergency Evacuation Plan should be kept in the secure information box?', 'Three copies.', ['Three copies of the Building Emergency Evacuation Plan'], page=4)
    add('gmfrs_information_box.pdf', 'What is the minimum number of riser outlet keys recommended in the fire sector folder?', 'At least 2, ideally 4.', ['Riser Outlet Keys', 'ideally 4', 'least 2'], page=11)
    add('ofgem_statement_of_intent.docx', 'On what date did GBIS Flex end according to this statement of intent template?', '31 March 2026', ['GBIS Flex', '31 March 2026'], paragraph=16)
    add('ofgem_statement_of_intent.docx', 'What information about qualifying schemes must be included for Route 2 Proxy 5?', 'The names and a brief description of the relevant schemes.', ['Proxy 5', 'names', 'brief description'], paragraph=10)
    add('ofgem_statement_of_intent.docx', 'What legal obligation does ECO4 place on energy suppliers?', 'Deliver energy efficiency measures to homes, including insulation and heating measures.', ['legal obligations', 'energy suppliers', 'insulation', 'heating'], paragraph=3)
    add('ofgem_medical_referral.docx', 'What age threshold is listed for older people vulnerable to cold under Route 2?', '65 years or older', ['65', 'older'], paragraph=19)
    add('ofgem_medical_referral.docx', 'From what kind of email address must a digital medical referral be sent?', 'A valid NHS email address.', ['valid NHS email address'], paragraph=31)
    add('ofgem_medical_referral.docx', 'Where applicable, what stamp authenticates a printed medical referral letter?', 'The GP surgery stamp.', ['GP surgery stamp'], paragraph=35)
    add('ofgem_householder_application.docx', 'What combined gross annual household income threshold applies to Route 1?', 'Less than 31,000 pounds.', ['gross annual', '31,000'], paragraph=24)
    add('ofgem_householder_application.docx', 'How many criteria must a household satisfy for Route 2?', 'At least two.', ['Route 2', 'at least two'], paragraph=21)
    add('ofgem_householder_application.docx', 'Does completing this householder application form guarantee eligibility for the schemes?', 'No; further assessments by the local authority and energy supplier are required.', ['does not guarantee eligibility', 'energy supplier'], paragraph=3)
    add('canterbury_fire_risk.docx', 'What test does the fire risk assessment ask about for landlord-supplied portable electrical appliances?', 'Portable appliance testing (PAT).', ['Landlord supplied portable appliances', 'PAT'], table=5, row=3)
    add('canterbury_fire_risk.docx', 'Which four groups of occupants are listed as being especially at risk from fire?', 'Disabled occupants, children, elderly occupants and other vulnerable adults.', ['Physically Disabled occupants', 'Children', 'Elderly', 'Other vulnerable adults'], table=3)
    add('canterbury_fire_risk.docx', 'Which action-plan columns record the responsible person, deadline and completion sign-off?', 'Who; When; Sign/date when completed.', ['Who', 'When', 'Sign/date when completed'], table=19, row=0)
    add('eia_electricity_prices.xlsx', 'What was the Total Electric Industry residential electricity price in 2024?', '16.48 cents per kilowatthour', ['16.48'], sheet='epa_02_04', cell='B15', value=16.48, unit='Cents per kilowatthour')
    add('eia_electricity_prices.xlsx', 'What was the Competitive Service Providers industrial electricity price in 2024?', '9.84 cents per kilowatthour', ['9.84'], sheet='epa_02_04', cell='D39', value=9.84, unit='Cents per kilowatthour')
    add('eia_electricity_prices.xlsx', 'What was the Delivery-Only Providers transportation electricity price in 2024?', '4.43 cents per kilowatthour', ['4.43'], sheet='epa_02_04', cell='E63', value=4.43, unit='Cents per kilowatthour')
    add('eia_industry_summary.xlsx', 'What was total coal net generation across all sectors in 2024?', '652,156 thousand megawatthours', ['652156'], sheet='epa_01_01', cell='C7', value=652156, unit='Thousand Megawatthours')
    add('eia_industry_summary.xlsx', 'How much natural gas was consumed for electricity generation in 2024 across all sectors?', '13,754,749 thousand Mcf', ['13754749'], sheet='epa_01_01', cell='C31', value=13754749, unit='Thousand Mcf')
    add('eia_industry_summary.xlsx', 'What was residential electricity sales revenue in 2024?', '244,367 million dollars', ['244367'], sheet='epa_01_01', cell='E48', value=244367, unit='Million Dollars')
    add('ofgem_declaration_template.xlsx', 'Which Sheet1 column stores the property postcode?', 'Column I, Post_Code.', ['Post_Code'], sheet='Sheet1', cell='I1', value='Post_Code')
    add('ofgem_declaration_template.xlsx', 'Which Sheet1 field records the Route 4 application number?', 'Column F, Route_4_Application_Number.', ['Route_4_Application_Number'], sheet='Sheet1', cell='F1', value='Route_4_Application_Number')
    add('ofgem_declaration_template.xlsx', 'Which Sheet1 field records the statement of intent publication date?', 'Column M, Date_Of_Statement_Of_Intent_Publication.', ['Date_Of_Statement_Of_Intent_Publication'], sheet='Sheet1', cell='M1', value='Date_Of_Statement_Of_Intent_Publication')
    return rows


def independent_source_audit(rows):
    from pypdf import PdfReader
    from docx import Document
    from openpyxl import load_workbook
    loaded = {}
    audits = []
    for sample in rows:
        name = sample['file']; loc = sample['native_locator']; path = FOLDER/'assets'/name
        if name not in loaded:
            loaded[name] = (PdfReader(path) if path.suffix == '.pdf' else
                            Document(path) if path.suffix == '.docx' else
                            load_workbook(path, read_only=True, data_only=True))
        source = loaded[name]
        if 'page' in loc: text = source.pages[loc['page']-1].extract_text()
        elif 'paragraph' in loc: text = source.paragraphs[loc['paragraph']].text
        elif 'table' in loc:
            table = source.tables[loc['table']]
            selected = [table.rows[loc['row']]] if 'row' in loc else table.rows
            text = '\n'.join(' | '.join(cell.text for cell in row.cells) for row in selected)
        else:
            value = source[loc['sheet']][loc['cell']].value
            text = str(value)
            if isinstance(loc['value'], str): equal = str(value).strip() == loc['value'].strip()
            else: equal = value == loc['value']
            if not equal: raise ValueError(f"Independent cell mismatch: {sample['id']}: {value!r}")
        valid = all(norm(anchor) in norm(text) for anchor in sample['anchors'])
        audits.append(dict(id=sample['id'], valid=valid, native_locator=loc, source_excerpt=text))
    for source in loaded.values():
        if hasattr(source, 'close'): source.close()
    return audits


def support(item, sample):
    if item.get('document_name') != sample['file'] or item.get('evidence_role') == 'navigation': return False
    text = item.get('text', '')
    if not all(norm(a) in norm(text) for a in sample['anchors']): return False
    loc = sample['native_locator']
    if 'cell' in loc:
        # The actual native cell and its value must occur on the same row.
        # Header questions deliberately use row 1; blank forms are not data rows.
        if f"sheet={loc['sheet']}" not in text: return False
        return any(re.search(r'(?<![A-Z0-9])'+re.escape(loc['cell'])+r'(?:\[[^\]]*\])?=', line)
                   and all(norm(a) in norm(line) for a in sample['anchors'])
                   for line in text.splitlines() if '[ROW' in line)
    if 'page' in loc:
        refs = ' '.join(item.get('source_refs', []))
        # Packing may remove source_refs, while keeping native BLOCK locators.
        return bool(re.search(r'pdf:page:'+str(loc['page'])+r'(?!\d)', refs+' '+text))
    if 'table' in loc:
        number = loc['table']+1  # python-docx tables are zero-based.
        return bool(re.search(r'word:(?:table:'+str(number)+r'(?!\d)|t'+str(number)+r':)', text))
    return True


def support_set(items, sample):
    if 'table' not in sample['native_locator']:
        return any(support(item,sample) for item in items)
    # A list answer may require several rows from the same native Word table.
    # Do not falsely require all four occupants to occur in one row/window.
    return all(any(support(item,dict(sample,anchors=[anchor])) for item in items)
               for anchor in sample['anchors'])


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    rows = cases()
    sources = json.loads((FOLDER/'download_manifest.json').read_text(encoding='utf-8'))
    for source in sources:
        if digest(FOLDER/source['path']) != source['sha256']: raise ValueError('Source hash changed')
    audits = independent_source_audit(rows)
    (OUTPUT/'source_gold_audit.json').write_text(json.dumps(audits, ensure_ascii=False, indent=2), encoding='utf-8')
    bad = [a['id'] for a in audits if not a['valid']]
    if bad: raise ValueError('Independent native gold failed before retrieval: '+str(bad))
    frozen = FOLDER/'english_questions.json'
    encoded = json.dumps(rows, ensure_ascii=False, indent=2)+'\n'
    if frozen.exists() and frozen.read_text(encoding='utf-8') != encoded:
        raise ValueError('Frozen question list differs; create a new version')
    if not frozen.exists():
        frozen.write_text(encoded, encoding='utf-8')
    old_hashes = set()
    for folder in [ROOT/'evaluation/public_v1/assets', ROOT/'evaluation/independent_native_20260917/assets', ROOT/'runtime/public_file_smoke_20260908/originals']:
        if folder.exists():
            old_hashes.update(digest(p) for p in folder.iterdir() if p.is_file())
    overlap = [s['file'] for s in sources if s['sha256'] in old_hashes]
    if overlap: raise ValueError('Previously tested source hash overlap: '+str(overlap))
    for key in ['CUSTOMER_DOCUMENT_OCR_ENABLED', 'CUSTOMER_ATTACHMENT_SEMANTIC_RERANK']: os.environ[key]='0'
    os.environ['RAG_HYBRID_ENABLED']=os.getenv('PUBLIC_EVAL_HYBRID_ENABLED','0')
    from backend.documents.customer_sessions import add_files, bind_session_owner, delete_session, get_session, retrieve
    from backend.sales.context_engine import optimise_evidence_context, validate_packed_evidence
    from backend.app import compact_grounded_payload_for_generation
    from transformers import AutoTokenizer
    tokenizer = AutoTokenizer.from_pretrained(os.getenv('PUBLIC_EVAL_TOKENIZER_PATH',str(ROOT/'models/Qwen3-VL-8B-Instruct')), local_files_only=True)
    sessions = {}; canonical = {}; results = []; owner = 'fresh-en-native-cpu'
    try:
        for sample in rows:
            started = time.perf_counter(); name = sample['file']
            result = dict(id=sample['id'], file=name, question=sample['question'], expected_answer=sample['expected_answer'])
            try:
                if name not in sessions:
                    sid = add_files([(name,(FOLDER/'assets'/name).read_bytes())], owner_id=owner)['session_id']
                    sessions[name]=sid
                    doc = get_session(sid, owner_id=owner).documents[0]
                    canonical[name] = dict(document_id=doc.document_id, parser=doc.parser, chunks=doc.chunks)
                with bind_session_owner(owner):
                    recalled = retrieve(sessions[name], sample['question'], max_chunks=24, max_text_tokens=5000)
                ranked, context_audit = optimise_evidence_context(sample['question'], recalled['evidence'])
                ranks = [i+1 for i in range(len(ranked)) if support_set(ranked[:i+1], sample)]
                payload, packing_audit = compact_grounded_payload_for_generation(
                    {'customer_question':sample['question'], 'evidence':ranked}, tokenizer,
                    max_prompt_tokens=int(os.getenv('PUBLIC_EVAL_PROMPT_BUDGET','1800')), system_prompt='Answer only from evidence.')
                packed = json.loads(payload)['evidence']
                # Score compacted text, but retain the exact source backpointer.
                # Never restore text removed by packing when checking coverage.
                originals = {item['evidence_id']:item for item in ranked}
                scored_packed = [dict(originals.get(item['evidence_id'], {}),**item) for item in packed]
                result.update(hit_at_5=any(r <= 5 for r in ranks), reciprocal_rank=1/min(ranks) if ranks else 0,
                              packed_support=support_set(scored_packed,sample),
                              unit_retained=not sample['native_locator'].get('unit') or norm(sample['native_locator']['unit']) in norm('\n'.join(item.get('text','') for item in packed)),
                              canonical_support=support_set([dict(c, document_name=name) for c in canonical[name]['chunks']],sample),
                              ranked_evidence=ranked, packed_evidence=packed,
                              retrieval_audit=recalled['input_snapshot'], context_audit=context_audit, packing_audit=packing_audit,
                              integrity=validate_packed_evidence(ranked,packed))
            except Exception as exc:
                result.update(hit_at_5=False, packed_support=False, reciprocal_rank=0, error={'type':type(exc).__name__,'message':str(exc)})
            result['elapsed_seconds']=round(time.perf_counter()-started,3)
            results.append(result)
            print(json.dumps({k:result.get(k) for k in ['id','file','hit_at_5','packed_support','canonical_support','error']},ensure_ascii=True),flush=True)
    finally:
        for sid in sessions.values(): delete_session(sid, owner_id=owner)
    groups = {'all':results}
    for ext in ['.pdf','.docx','.xlsx']: groups[ext]=[r for r in results if Path(r['file']).suffix==ext]
    metrics = {group:dict(cases=len(items), hits=sum(r['hit_at_5'] for r in items), recall_at_5=sum(r['hit_at_5'] for r in items)/len(items),
                         mrr=sum(r['reciprocal_rank'] for r in items)/len(items), packed_support=sum(r['packed_support'] for r in items))
               for group,items in groups.items()}
    report = dict(track='CPU_single_file_native_preprocessing_lexical_retrieval_context_not_generation',
                  scoring_version='native_source_complete_support_set_v2',
                  scoring_definition='Complete source-grounded answer support in Top5; native Word list answers may span several windows from the exact same table.',
                  model_generation_tested=False, planner_tested=False, semantic_reranker_tested=False, ocr_tested=False,
                  independent_gold_valid=len(audits), previously_tested_sha_overlap=overlap,
                  question_sha256=digest(frozen), source_files=len(sources), metrics=metrics,
                  failed_cases=[r['id'] for r in results if not r['hit_at_5']], details=results)
    (OUTPUT/'results.json').write_text(json.dumps(report, ensure_ascii=False, indent=2),encoding='utf-8')
    (OUTPUT/'canonical_snapshot.json').write_text(json.dumps(canonical,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(metrics),flush=True)


if __name__ == '__main__': main()
