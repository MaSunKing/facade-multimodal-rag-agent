"""Synthetic, local HTTP intake checks; deliberately never request generation."""
from __future__ import annotations

import argparse
import hashlib
from io import BytesIO
import json
import os
from pathlib import Path
import sys
import time

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ.setdefault("CUSTOMER_DOCUMENT_OCR_ENABLED", "0")
os.environ.setdefault("CUSTOMER_ATTACHMENT_SEMANTIC_RERANK", "0")
os.environ.setdefault("RAG_HYBRID_ENABLED", "0")


def fixtures():
    import fitz
    from PIL import Image, ImageDraw
    from docx import Document
    from openpyxl import Workbook
    folder = ROOT / "runtime/public_smoke/documents"
    folder.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (640, 360), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((80, 90, 260, 230), outline="black", width=3)
    draw.rectangle((380, 90, 560, 230), outline="black", width=3)
    draw.line((260, 160, 380, 160), fill="black", width=3)
    draw.text((90, 40), "SYNTHETIC DEMO - not a construction detail", fill="black")
    draw.text((100, 140), "Panel", fill="black")
    draw.text((400, 140), "Support", fill="black")
    png = BytesIO(); image.save(png, format="PNG")
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((50, 50), "SYNTHETIC DemoPanel-X. Thickness: 20 mm. Not a real specification.")
    page.insert_image(fitz.Rect(50, 100, 530, 370), stream=png.getvalue())
    page = pdf.new_page()
    page.insert_text((50, 50), "Revision note: another source states 18 mm. Verify versions before use.")
    pdf_bytes = pdf.tobytes(); pdf.close()
    word = Document()
    word.add_heading("Synthetic installation checklist", 0)
    word.add_paragraph("DemoPanel-X: verify thickness, substrate and fixing design. Not engineering advice.")
    table = word.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Item"; table.rows[0].cells[1].text = "Requirement"
    cells = table.add_row().cells; cells[0].text = "Thickness"; cells[1].text = "20 mm"
    word.add_picture(BytesIO(png.getvalue()))
    docx = BytesIO(); word.save(docx)
    workbook = Workbook(); sheet = workbook.active; sheet.title = "Current"
    sheet.append(["Product", "Thickness", "Unit", "Source version"])
    sheet.append(["DemoPanel-X", 20, "mm", "A"])
    other = workbook.create_sheet("Other source")
    other.append(["Product", "Thickness", "Unit", "Source version"])
    other.append(["DemoPanel-X", 18, "mm", "B"])
    xlsx = BytesIO(); workbook.save(xlsx); workbook.close()
    files = {"sample_product.pdf": pdf_bytes, "sample_checklist.docx": docx.getvalue(),
             "sample_spec.xlsx": xlsx.getvalue(), "sample_drawing.png": png.getvalue()}
    for name, content in files.items():
        (folder / name).write_bytes(content)
    return files


def run(client, files):
    checks = []
    def check(name, condition):
        checks.append({"check": name, "passed": bool(condition)})
        if not condition:
            raise AssertionError(name)
    started = time.perf_counter()
    owner = {"x-facade-client-id": "synthetic_public_owner_123456"}
    foreign = {"x-facade-client-id": "synthetic_foreign_owner_123456"}
    check("health_live", client.get("/health/live").status_code == 200)
    health = client.get("/health/ready")
    check("health_ready_endpoint", health.status_code == 200)
    check("legacy_finance_api_absent", not any("/finance" in p for p in client.get("/openapi.json").json()["paths"]))
    check("upload_requires_owner", client.post("/api/copilot/documents",
        files=[("files", ("sample.txt", b"synthetic"))]).status_code == 400)
    uploaded = client.post("/api/copilot/documents", headers=owner,
        files=[("files", (name, data)) for name, data in files.items()])
    check("four_native_formats_upload", uploaded.status_code == 200)
    summary = uploaded.json(); sid = summary["session_id"]
    try:
        documents = summary["documents"]
        check("all_four_files_preserved", len(documents) == 4)
        check("text_formats_have_chunks", all(d["chunk_count"] > 0 for d in documents if d["file_name"] != "sample_drawing.png"))
        check("pdf_visual_retained", any(d["file_name"].endswith(".pdf") and d["ready_visual_count"] > 0 for d in documents))
        check("word_visual_retained", any(d["file_name"].endswith(".docx") and d["ready_visual_count"] > 0 for d in documents))
        check("owner_can_read", client.get(f"/api/copilot/documents/{sid}", headers=owner).status_code == 200)
        check("foreign_cannot_read", client.get(f"/api/copilot/documents/{sid}", headers=foreign).status_code == 404)
        check("foreign_cannot_append", client.post("/api/copilot/documents", headers=foreign,
            data={"session_id": sid}, files=[("files", ("extra.txt", b"synthetic"))]).status_code == 404)
        check("foreign_cannot_delete", client.delete(f"/api/copilot/documents/{sid}", headers=foreign).status_code == 404)
        from backend.document_parsing.file_ingestion import ingest_uploaded_file
        parsed = ingest_uploaded_file("sample_drawing.png", files["sample_drawing.png"])
        visual_id = parsed.intermediate.visual_assets[0].visual_id
        doc_id = next(d["document_id"] for d in documents if d["file_name"] == "sample_drawing.png")
        visual_url = f"/api/copilot/documents/{sid}/visual/{doc_id}/{visual_id}"
        served = client.get(visual_url, headers=owner)
        check("original_image_bytes_match", served.status_code == 200 and hashlib.sha256(served.content).digest() == hashlib.sha256(files["sample_drawing.png"]).digest())
        check("foreign_cannot_read_image", client.get(visual_url, headers=foreign).status_code == 404)
        # Independently audit native parser structures, not just HTTP statuses.
        workbook = ingest_uploaded_file("sample_spec.xlsx", files["sample_spec.xlsx"])
        check("excel_two_sheets_preserved", len({t.sheet_name for t in workbook.intermediate.tables}) == 2)
        from backend.documents.evidence_v2 import EvidenceDocumentV2, convert_intermediate_to_v2
        first = convert_intermediate_to_v2(workbook.intermediate)
        second = convert_intermediate_to_v2(workbook.intermediate)
        EvidenceDocumentV2.model_validate(first.model_dump())
        check("evidence_schema_valid", True)
        check("snapshot_and_ids_stable", first.model_dump() == second.model_dump())
        check("owner_can_delete", client.delete(f"/api/copilot/documents/{sid}", headers=owner).status_code == 200)
        check("deleted_session_unavailable", client.get(f"/api/copilot/documents/{sid}", headers=owner).status_code == 404)
    finally:
        client.delete(f"/api/copilot/documents/{sid}", headers=owner)
    return {"fixture_provenance": "project_authored_synthetic", "model_generation_tested": False,
            "check_count": len(checks), "passed": all(c["passed"] for c in checks), "checks": checks,
            "elapsed_seconds": round(time.perf_counter() - started, 3)}


def main():
    parser = argparse.ArgumentParser(); parser.add_argument("--base-url")
    args = parser.parse_args(); files = fixtures()
    if args.base_url:
        from urllib.parse import urlparse
        if urlparse(args.base_url).hostname not in {"localhost", "127.0.0.1", "::1"}:
            parser.error("This smoke test is restricted to a local server.")
        import httpx
        with httpx.Client(base_url=args.base_url, timeout=90) as client:
            report = run(client, files)
    else:
        from fastapi.testclient import TestClient
        from backend.app import app
        with TestClient(app) as client:
            report = run(client, files)
    output = ROOT / "runtime/public_smoke/http_report.json"
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    main()
