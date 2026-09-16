"""Private, non-persistent intake for finance data.

The intake boundary supports text, XLSX/XLS/CSV, PDF, Word, and image files.  It
has two outputs:

* an *intermediate document*, which faithfully records what the parser saw;
* a *standard financial document*, which contains only deterministic field
  mappings and keeps a pointer back to the original source cell or text.

No company file is written to disk and no LLM is called in this module.  That
makes parser failures, mapping failures, and any future model failures easy to
separate and audit.
"""

from __future__ import annotations

import csv
import hashlib
import math
import os
import re
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any, Iterable, Literal

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import range_to_tuple
from pydantic import BaseModel, Field
from backend.documents.visual_layout import compact_visual_metadata as _searchable_visual_metadata


MAX_EXCEL_BYTES = 5 * 1024 * 1024
# Crossing this threshold enables auditable large-workbook mode; it is not an
# extraction failure.  Canonical evidence remains lossless while model input is
# still emitted through separately bounded, source-linked context chunks.
LARGE_EXCEL_NONEMPTY_CELL_THRESHOLD = 100_000
# This is a resource/DoS guard, not a semantic document-size assumption.  A
# deployment with a different memory budget can override it without changing
# parsing behavior or adding workbook-specific exceptions.
try:
    MAX_EXCEL_NONEMPTY_CELLS = max(
        LARGE_EXCEL_NONEMPTY_CELL_THRESHOLD,
        int(os.getenv("DOCUMENT_EXCEL_MAX_NONEMPTY_CELLS", "500000")),
    )
except ValueError:
    MAX_EXCEL_NONEMPTY_CELLS = 500_000
try:
    MAX_EXCEL_GRID_CELLS = max(
        LARGE_EXCEL_NONEMPTY_CELL_THRESHOLD,
        int(os.getenv("DOCUMENT_EXCEL_MAX_GRID_CELLS", "500000")),
    )
except ValueError:
    MAX_EXCEL_GRID_CELLS = 500_000
MAX_TEXT_CHARACTERS = 8_000
MAX_MODEL_CONTEXT_CHARACTERS = 60_000
MAX_MODEL_CONTEXT_CHUNK_CHARACTERS = 12_000
MAX_REPEATED_CHUNK_HEADER_CHARACTERS = 4_000
MAX_DERIVED_TEXT_CHUNK_CHARACTERS = 1_600
MAX_EXCEL_VISUAL_ASSETS = 8
MAX_EXCEL_VISUAL_ASSET_BYTES = 4 * 1024 * 1024
NON_INGESTION_SHEET_TOKENS = ("示例", "填写说明", "字段字典", "example", "instruction", "dictionary")


class FinanceIntakeError(ValueError):
    """Raised when an uploaded input cannot be safely processed."""


class SourcePointer(BaseModel):
    source_type: Literal["text", "excel", "pdf", "word", "image", "html", "xml"]
    file_name: str | None = None
    sheet_name: str | None = None
    page_number: int | None = None
    cell: str | None = None
    row_index: int | None = None
    column_index: int | None = None
    section_id: str | None = None
    section_title: str | None = None
    text_span: tuple[int, int] | None = None
    parser: str | None = None
    extraction_confidence: float | None = Field(default=None, ge=0, le=1)
    # Native Office/PDF parsers do not always expose physical coordinates.
    # OCR can later add a page-image bounding box without changing consumers.
    bounding_box: tuple[float, float, float, float] | None = None


class NumericCandidate(BaseModel):
    """A reversible numeric surface form, not a financial fact.

    ``normalized_value`` only applies an explicit unit multiplier (for example
    ``万``).  It never assumes a currency, financial metric, sign convention,
    or report unit; those are semantic decisions.
    """

    raw_text: str
    numeric_value: float
    unit_token: str | None = None
    multiplier: float = 1.0
    normalized_value: float | None = None
    requires_semantic_interpretation: bool = True


class EvidenceBlock(BaseModel):
    """A source-addressable non-tabular or tabular document fragment."""

    block_id: str
    kind: Literal[
        "worksheet",
        "table",
        "paragraph",
        "page_text",
        "header",
        "footer",
        "comment",
        "revision",
        "image_metadata",
        "embedded_visual",
        "derived_text",
    ]
    source: SourcePointer
    text: str | None = None
    table_id: str | None = None
    numeric_candidates: list[NumericCandidate] = Field(default_factory=list)
    metadata: dict[str, Any] = Field(default_factory=dict)


def _annotate_large_tabular_document(blocks: list[EvidenceBlock], populated_cell_count: int) -> None:
    """Record capacity and coverage without changing extracted cell evidence."""

    large_workbook_mode = populated_cell_count > LARGE_EXCEL_NONEMPTY_CELL_THRESHOLD
    for block in blocks:
        if block.kind != "worksheet":
            continue
        block.metadata.update(
            {
                "workbook_populated_cell_count": populated_cell_count,
                "large_workbook_mode": large_workbook_mode,
                "extraction_coverage_complete": True,
                "hard_populated_cell_limit": MAX_EXCEL_NONEMPTY_CELLS,
                "model_context_chunk_character_limit": MAX_MODEL_CONTEXT_CHUNK_CHARACTERS,
            }
        )


class VisualAsset(BaseModel):
    """A visual companion to the tabular Evidence JSON.

    The binary image is deliberately excluded from API JSON.  It remains only
    in the in-memory request object long enough for the local vision model to
    inspect it, while the returned metadata stays auditable and compact.
    """

    visual_id: str
    kind: Literal["embedded_image", "chart_preview", "unrendered_chart", "document_page"]
    source: SourcePointer
    media_type: str | None = None
    width_px: int | None = Field(default=None, ge=1)
    height_px: int | None = Field(default=None, ge=1)
    sha256: str | None = None
    delivery_status: Literal["ready_for_vision", "metadata_only", "rendering_unavailable"]
    metadata: dict[str, Any] = Field(default_factory=dict)
    image_bytes: bytes | None = Field(default=None, exclude=True, repr=False)


class VisualObservation(BaseModel):
    """A non-numeric, reviewable observation from a local vision model."""

    visual_id: str
    status: Literal["candidate_ready", "unavailable", "failed"]
    chart_type: str | None = Field(default=None, max_length=80)
    title: str | None = Field(default=None, max_length=240)
    findings: list[str] = Field(default_factory=list, max_length=6)
    confidence: float = Field(default=0.0, ge=0, le=1)
    requires_confirmation: bool = True
    source_of_truth: Literal["supplemental_visual_only"] = "supplemental_visual_only"
    message: str | None = Field(default=None, max_length=300)


class VisionTableCandidate(BaseModel):
    """A review-only table transcription proposed by the local vision model.

    Values here are deliberately *not* ``IntermediateCell`` instances.  They
    have no deterministic OCR coordinate and must never become confirmed
    business facts or final output until the customer confirms them.
    """

    title: str | None = Field(default=None, max_length=160)
    headers: list[str] = Field(default_factory=list, max_length=16)
    rows: list[list[str]] = Field(default_factory=list, max_length=80)


class VisionDocumentCandidate(BaseModel):
    """Bounded, auditable visual transcription for images and scanned PDFs."""

    visual_id: str
    status: Literal["candidate_ready", "unavailable", "failed"]
    document_type: str | None = Field(default=None, max_length=80)
    title: str | None = Field(default=None, max_length=240)
    text_blocks: list[str] = Field(default_factory=list, max_length=24)
    tables: list[VisionTableCandidate] = Field(default_factory=list, max_length=4)
    confidence: float = Field(default=0.0, ge=0, le=1)
    recognizer: Literal["qwen3_vl", "pp_structure_v3"] = "qwen3_vl"
    requires_confirmation: bool = True
    source_of_truth: Literal["vision_candidate_only"] = "vision_candidate_only"
    message: str | None = Field(default=None, max_length=300)


class IntermediateCell(BaseModel):
    coordinate: str
    row_index: int
    column_index: int
    value: str | int | float | bool | None = None
    display_text: str | None = None
    formula: str | None = None
    cached_value: str | int | float | bool | None = None
    numeric_candidates: list[NumericCandidate] = Field(default_factory=list)


class IntermediateTable(BaseModel):
    table_id: str
    # A sheet can contain several disconnected logical tables.  ``table_id``
    # identifies one block; ``parent_table_id`` groups sibling blocks from the
    # same source sheet for scope comparisons.
    parent_table_id: str | None = None
    section_id: str | None = None
    title: str | None = None
    title_source: SourcePointer | None = None
    sheet_name: str | None = None
    page_number: int | None = None
    sheet_state: str | None = None
    range: str | None = None
    parser: str | None = None
    extraction_confidence: float | None = Field(default=None, ge=0, le=1)
    header_row_index: int | None = Field(default=None, ge=1)
    header_range: str | None = None
    scope: Literal["includes_prior_periods", "current_period_only", "unknown"] = "unknown"
    scope_confidence: float = Field(default=0.0, ge=0, le=1)
    scope_evidence: SourcePointer | None = None
    headers: list[str] = Field(default_factory=list)
    rows: list[list[IntermediateCell]] = Field(default_factory=list)
    merged_ranges: list[str] = Field(default_factory=list)


class LanguageCandidate(BaseModel):
    """A conservative language/script hint; never a translation or fact."""

    code: str = Field(min_length=2, max_length=24)
    script: str = Field(min_length=2, max_length=24)
    confidence: float = Field(ge=0, le=1)


class UnitCandidate(BaseModel):
    """An explicitly observed currency or scale label with its source."""

    kind: Literal["currency", "scale"]
    raw_text: str = Field(min_length=1, max_length=80)
    currency_code: str | None = Field(default=None, max_length=8)
    scale_multiplier: float | None = Field(default=None, gt=0)
    source: SourcePointer
    confidence: float = Field(ge=0, le=1)
    requires_confirmation: bool = True


class StructureProfile(BaseModel):
    """Generic structural description of one source-addressable evidence area."""

    profile_id: str = Field(min_length=1, max_length=200)
    kind: Literal["document", "worksheet", "table", "block", "visual"]
    role: Literal["structured_table", "key_value", "narrative_note", "mixed", "visual_only", "unknown"]
    source: SourcePointer
    confidence: float = Field(ge=0, le=1)
    signals: list[str] = Field(default_factory=list, max_length=12)
    language_candidates: list[LanguageCandidate] = Field(default_factory=list, max_length=4)
    unit_candidates: list[UnitCandidate] = Field(default_factory=list, max_length=16)


class ModelContextChunk(BaseModel):
    """A complete, source-linked model-read chunk; no canonical evidence is dropped."""

    chunk_id: str = Field(min_length=1, max_length=200)
    kind: Literal["document_index", "evidence"]
    text: str = Field(min_length=1)
    character_count: int = Field(ge=1)
    source_refs: list[str] = Field(default_factory=list, max_length=240)


class IntermediateDocument(BaseModel):
    """Request-scoped evidence before financial semantic interpretation.

    ``tables`` preserves calculation-grade grids.  ``blocks`` preserves text,
    sheet/page structure and provenance.  ``model_context`` is merely a
    bounded rendering for a model, never the canonical evidence.
    """

    evidence_schema_version: str = "evidence-document-v1"
    document_id: str | None = None
    file_sha256: str | None = None
    file_size_bytes: int | None = Field(default=None, ge=0)
    source_type: Literal["text", "excel", "pdf", "word", "image", "html", "xml"]
    file_name: str | None = None
    parser: str
    tables: list[IntermediateTable]
    raw_text: str | None = None
    blocks: list[EvidenceBlock] = Field(default_factory=list)
    visual_assets: list[VisualAsset] = Field(default_factory=list)
    structure_profiles: list[StructureProfile] = Field(default_factory=list)
    model_context_chunks: list[ModelContextChunk] = Field(default_factory=list)
    model_context: str | None = None


class StandardFinancialFact(BaseModel):
    statement_type: Literal["income_statement", "balance_sheet", "operating_data"]
    standard_item_code: str
    standard_item_name: str
    original_item_name: str
    period: str | None = None
    value: float | None = None
    currency: str = "CNY"
    unit: Literal["yuan", "unknown"] = "yuan"
    source: SourcePointer
    mapping_method: Literal["rule", "semantic_candidate", "customer_confirmed"] = "rule"
    mapping_confidence: float = Field(ge=0, le=1)
    derived_from_codes: list[str] = Field(default_factory=list, max_length=12)


class StandardFinancialDocument(BaseModel):
    company_name: str | None = None
    facts: list[StandardFinancialFact] = Field(default_factory=list)


class ValidationIssue(BaseModel):
    severity: Literal["error", "warning"]
    code: str
    message: str
    source: SourcePointer | None = None


class IntakeResult(BaseModel):
    intermediate: IntermediateDocument
    standard: StandardFinancialDocument
    validation: list[ValidationIssue]
    visual_observations: list[VisualObservation] = Field(default_factory=list)
    vision_document_candidates: list[VisionDocumentCandidate] = Field(default_factory=list)


@dataclass(frozen=True)
class MetricDefinition:
    code: str
    name: str
    statement_type: Literal["income_statement", "balance_sheet", "operating_data"]
    aliases: tuple[str, ...]


class SemanticLabelMapping(BaseModel):
    """A local-model label interpretation, never an independently generated fact.

    ``target`` may be one of the financial metric codes or a structural role
    such as ``__PERIOD__``.  The program still reads the amount, period and
    source coordinate from the uploaded table itself.
    """

    table_id: str | None = Field(default=None, max_length=160)
    label: str = Field(min_length=1, max_length=160)
    target: str = Field(min_length=1, max_length=80)
    confidence: float = Field(ge=0, le=1)
    rationale: str | None = Field(default=None, max_length=240)


class SemanticTableContext(BaseModel):
    """Table-level semantic context stated by the local model from the upload."""

    table_id: str = Field(min_length=1, max_length=160)
    period_year: str | None = Field(default=None, pattern=r"^20\d{2}$")


METRICS: tuple[MetricDefinition, ...] = (
    MetricDefinition("IS_REVENUE", "营业收入", "income_statement", ("营业总收入", "主营业务收入", "营业收入", "营业收入净额", "收入", "revenue")),
    MetricDefinition("IS_COST_OF_SALES", "营业成本", "income_statement", ("主营业务成本", "销售成本", "营业成本", "成本", "cost_of_sales")),
    MetricDefinition("IS_SALES_EXPENSE", "销售费用", "income_statement", ("销售费用", "市场费用", "sales_expense")),
    MetricDefinition("IS_ADMINISTRATIVE_EXPENSE", "管理费用", "income_statement", ("管理费用", "行政费用", "administrative_expense")),
    MetricDefinition("OD_BUDGET_REVENUE", "收入预算", "operating_data", ("营业收入预算", "收入预算", "预算收入", "budget_revenue")),
    MetricDefinition("OD_BUDGET_OPERATING_PROFIT", "营业利润预算", "operating_data", ("营业利润预算", "经营利润预算", "budget_operating_profit")),
    MetricDefinition("BS_ACCOUNTS_RECEIVABLE", "应收账款", "balance_sheet", ("期末应收账款", "应收账款", "应收款", "accounts_receivable")),
    MetricDefinition("BS_CASH_BALANCE", "期末现金余额", "balance_sheet", ("期末现金余额", "现金余额", "货币资金", "cash_balance")),
)


def _normalised_label(value: str) -> str:
    return re.sub(r"[\s_\-（）()：:、]", "", value).lower()


_METRIC_BY_NORMALISED_ALIAS = {
    alias: metric
    for metric in METRICS
    for alias in (_normalised_label(item) for item in metric.aliases)
}
_METRIC_BY_CODE = {metric.code: metric for metric in METRICS}
_PERIOD_HEADER_LABELS = {"period", "月份", "期间", "报告期间", "日期", "date", "month"}
_ITEM_HEADER_LABELS = {"项目", "科目", "item", "项目名称", "科目名称"}
_CURRENT_AMOUNT_HEADER_LABELS = {"本期金额", "本年累计数", "本期数", "期末余额", "amount", "金额", "发生额"}
_PRIOR_AMOUNT_HEADER_LABELS = {"上期金额", "上年累计数", "上期数", "期初余额", "prior_amount", "上年金额"}
_STRUCTURAL_TARGETS = {"__PERIOD__", "__ITEM__", "__CURRENT_AMOUNT__", "__PRIOR_AMOUNT__"}


def _serialise_cell_value(value: Any) -> str | int | float | bool | None:
    if value is None:
        return None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _display(value: Any) -> str | None:
    serialised = _serialise_cell_value(value)
    return str(serialised) if serialised is not None else None


_SURFACE_NUMBER = re.compile(
    r"(?<![\w.])(?P<raw>\(?[-+]?\d+(?:,\d{3})*(?:\.\d+)?\)?)(?:\s*(?P<unit>亿元|万元|元|亿|万))?"
)
_SURFACE_MULTIPLIERS = {"万": 10_000.0, "万元": 10_000.0, "亿": 100_000_000.0, "亿元": 100_000_000.0}


def numeric_candidates_for_text(value: Any) -> list[NumericCandidate]:
    """Keep explicit numeric surface forms without turning them into facts."""

    text = _display(value)
    if not text:
        return []
    candidates: list[NumericCandidate] = []
    for match in _SURFACE_NUMBER.finditer(text):
        raw = match.group("raw")
        normalized = raw.replace(",", "")
        if normalized.startswith("(") and normalized.endswith(")"):
            normalized = f"-{normalized[1:-1]}"
        try:
            numeric_value = float(normalized)
        except ValueError:
            continue
        if not math.isfinite(numeric_value):
            continue
        unit = match.group("unit")
        multiplier = _SURFACE_MULTIPLIERS.get(unit or "", 1.0)
        candidates.append(
            NumericCandidate(
                raw_text=match.group(0),
                numeric_value=numeric_value,
                unit_token=unit,
                multiplier=multiplier,
                normalized_value=numeric_value * multiplier if unit else None,
            )
        )
    return candidates


def evidence_block(
    *,
    block_id: str,
    kind: Literal[
        "worksheet", "table", "paragraph", "page_text", "header", "footer", "comment", "revision", "image_metadata", "embedded_visual", "derived_text"
    ],
    source: SourcePointer,
    text: str | None = None,
    table_id: str | None = None,
    metadata: dict[str, Any] | None = None,
) -> EvidenceBlock:
    """Create a consistently provenance-tagged evidence block."""

    return EvidenceBlock(
        block_id=block_id,
        kind=kind,
        source=source,
        text=text,
        table_id=table_id,
        numeric_candidates=numeric_candidates_for_text(text),
        metadata=metadata or {},
    )


def _source_ref(source: SourcePointer) -> str:
    """Create a stable, human-readable locator without inventing a location."""

    parts = [source.source_type]
    if source.sheet_name:
        parts.append(f"sheet={source.sheet_name}")
    if source.page_number:
        parts.append(f"page={source.page_number}")
    if source.cell:
        parts.append(f"cell={source.cell}")
    if source.section_id:
        parts.append(f"section={source.section_id}")
    if source.text_span:
        parts.append(f"span={source.text_span[0]}:{source.text_span[1]}")
    return ";".join(parts)


_LANGUAGE_SCRIPTS: tuple[tuple[str, str, str, re.Pattern[str]], ...] = (
    ("zh", "Hani", "CJK", re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff]")),
    ("ja", "Jpan", "Japanese", re.compile(r"[\u3040-\u30ff]")),
    ("ko", "Hang", "Korean", re.compile(r"[\uac00-\ud7af\u1100-\u11ff]")),
    ("ar", "Arab", "Arabic", re.compile(r"[\u0600-\u06ff\u0750-\u077f]")),
    ("he", "Hebr", "Hebrew", re.compile(r"[\u0590-\u05ff]")),
    ("und", "Deva", "Devanagari", re.compile(r"[\u0900-\u097f]")),
    ("th", "Thai", "Thai", re.compile(r"[\u0e00-\u0e7f]")),
    ("und", "Cyrl", "Cyrillic", re.compile(r"[\u0400-\u052f]")),
    ("und", "Grek", "Greek", re.compile(r"[\u0370-\u03ff]")),
    ("und", "Latn", "Latin", re.compile(r"[A-Za-zÀ-ÖØ-öø-ÿ]")),
)
_ENGLISH_WORDS = re.compile(r"\b(?:the|and|of|for|report|statement|total|amount|year|date|income|balance)\b", re.IGNORECASE)


def language_candidates_for_text(value: Any) -> list[LanguageCandidate]:
    """Return low-risk script hints while preserving the source text unchanged.

    Language identification is deliberately conservative: a Latin script is
    only called English when several ordinary English words are visible.  A
    script hint is still useful for routing and prompts, but is never used to
    translate, map a financial item, or select a reporting template.
    """

    text = _display(value) or ""
    if not text.strip():
        return []
    counts: list[tuple[str, str, str, int]] = []
    for code, script, label, pattern in _LANGUAGE_SCRIPTS:
        count = len(pattern.findall(text))
        if count:
            counts.append((code, script, label, count))
    if not counts:
        return [LanguageCandidate(code="und", script="Zyyy", confidence=0.2)]

    # Japanese normally contains both CJK ideographs and kana.  Treat the
    # combined writing system as Japanese rather than emitting conflicting
    # Chinese and Japanese guesses.
    japanese = next((item for item in counts if item[1] == "Jpan"), None)
    if japanese is not None:
        cjk = next((item for item in counts if item[1] == "Hani"), None)
        if cjk is not None:
            counts = [item for item in counts if item[1] not in {"Jpan", "Hani"}]
            counts.append(("ja", "Jpan", "Japanese", japanese[3] + cjk[3]))

    total = sum(item[3] for item in counts)
    candidates: list[LanguageCandidate] = []
    for code, script, _label, count in sorted(counts, key=lambda item: item[3], reverse=True)[:3]:
        resolved_code = "en" if script == "Latn" and _ENGLISH_WORDS.search(text) else code
        candidates.append(
            LanguageCandidate(
                code=resolved_code,
                script=script,
                confidence=round(min(0.99, max(0.2, count / max(total, 1))), 3),
            )
        )
    return candidates


_UNIT_PATTERNS: tuple[tuple[Literal["currency", "scale"], re.Pattern[str], str | None, float | None, float], ...] = (
    ("currency", re.compile(r"\b(?:USD|US\$|U\.S\.\s*dollars?)\b", re.IGNORECASE), "USD", None, 0.98),
    ("currency", re.compile(r"\bEUR\b|€", re.IGNORECASE), "EUR", None, 0.98),
    ("currency", re.compile(r"\bGBP\b|£", re.IGNORECASE), "GBP", None, 0.98),
    ("currency", re.compile(r"\bJPY\b", re.IGNORECASE), "JPY", None, 0.98),
    ("currency", re.compile(r"(?:人民币|人民幣|RMB|CNY|Chinese\s+Yuan)", re.IGNORECASE), "CNY", None, 0.98),
    # A bare dollar or yen symbol is retained as observed evidence, but never
    # assigned a currency code because the symbol is genuinely ambiguous.
    ("currency", re.compile(r"[$¥]"), None, None, 0.45),
    ("scale", re.compile(r"(?:in\s+)?thousands?\b", re.IGNORECASE), None, 1_000.0, 0.95),
    ("scale", re.compile(r"(?:in\s+)?millions?\b", re.IGNORECASE), None, 1_000_000.0, 0.95),
    ("scale", re.compile(r"(?:in\s+)?billions?\b", re.IGNORECASE), None, 1_000_000_000.0, 0.95),
    ("scale", re.compile(r"(?:万元|萬\s*元)"), None, 10_000.0, 0.98),
    ("scale", re.compile(r"(?:亿元|億\s*元)"), None, 100_000_000.0, 0.98),
)


def unit_candidates_for_text(value: Any, source: SourcePointer) -> list[UnitCandidate]:
    """Find only explicit unit labels and retain their original source.

    A candidate is an observation, not a conversion instruction.  The
    customer or later analysis still confirms whether a sheet-level label
    applies to a particular number.
    """

    text = _display(value) or ""
    candidates: list[UnitCandidate] = []
    seen: set[tuple[str, str, str | None, float | None]] = set()
    for kind, pattern, currency_code, scale_multiplier, confidence in _UNIT_PATTERNS:
        for match in pattern.finditer(text):
            raw_text = match.group(0)
            identity = (kind, raw_text.casefold(), currency_code, scale_multiplier)
            if identity in seen:
                continue
            seen.add(identity)
            candidates.append(
                UnitCandidate(
                    kind=kind,
                    raw_text=raw_text,
                    currency_code=currency_code,
                    scale_multiplier=scale_multiplier,
                    source=source,
                    confidence=confidence,
                )
            )
    return candidates


def _table_source(document: IntermediateDocument, table: IntermediateTable) -> SourcePointer:
    return SourcePointer(
        source_type=document.source_type,
        file_name=document.file_name,
        sheet_name=table.sheet_name,
        page_number=table.page_number,
        section_id=table.section_id or table.table_id,
        parser=table.parser or document.parser,
        extraction_confidence=table.extraction_confidence,
    )


def _deduplicate_unit_candidates(candidates: Iterable[UnitCandidate], *, limit: int = 16) -> list[UnitCandidate]:
    result: list[UnitCandidate] = []
    seen: set[tuple[str, str, str | None, float | None, str]] = set()
    for candidate in candidates:
        identity = (
            candidate.kind,
            candidate.raw_text.casefold(),
            candidate.currency_code,
            candidate.scale_multiplier,
            _source_ref(candidate.source),
        )
        if identity not in seen:
            seen.add(identity)
            result.append(candidate)
        if len(result) >= limit:
            break
    return result


def _table_structure_profile(document: IntermediateDocument, table: IntermediateTable) -> StructureProfile:
    populated = [cell for row in table.rows for cell in row if cell.display_text not in {None, ""} or cell.formula]
    row_count = len(table.rows)
    column_count = max((len(row) for row in table.rows), default=0)
    values = [cell.display_text or cell.formula or "" for cell in populated]
    numeric_count = sum(bool(cell.numeric_candidates) for cell in populated)
    long_text_count = sum(len(value) >= 160 for value in values)
    signals = [f"rows={row_count}", f"columns={column_count}", f"populated_cells={len(populated)}"]
    if table.merged_ranges:
        signals.append(f"merged_ranges={len(table.merged_ranges)}")
    if numeric_count:
        signals.append(f"numeric_surface_cells={numeric_count}")
    if long_text_count:
        signals.append(f"long_text_cells={long_text_count}")

    if not populated:
        role: Literal["structured_table", "key_value", "narrative_note", "mixed", "visual_only", "unknown"] = "unknown"
        confidence = 0.4
    elif column_count <= 2 and long_text_count and long_text_count >= max(1, len(populated) // 3):
        role, confidence = "narrative_note", 0.75
    elif column_count == 2 and row_count >= 2 and numeric_count <= len(populated) // 2:
        role, confidence = "key_value", 0.68
    elif column_count >= 2 and row_count >= 2:
        role, confidence = "structured_table", 0.82
    elif long_text_count:
        role, confidence = "narrative_note", 0.65
    else:
        role, confidence = "mixed", 0.55

    text_for_language = "\n".join(values[:300])
    unit_candidates: list[UnitCandidate] = []
    for cell in populated:
        cell_source = SourcePointer(
            source_type=document.source_type,
            file_name=document.file_name,
            sheet_name=table.sheet_name,
            page_number=table.page_number,
            cell=cell.coordinate,
            row_index=cell.row_index,
            column_index=cell.column_index,
            section_id=table.section_id or table.table_id,
            parser=table.parser or document.parser,
            extraction_confidence=table.extraction_confidence,
        )
        unit_candidates.extend(unit_candidates_for_text(cell.display_text or cell.formula, cell_source))
        if len(unit_candidates) >= 24:
            break
    return StructureProfile(
        profile_id=f"profile:{table.table_id}",
        kind="table",
        role=role,
        source=_table_source(document, table),
        confidence=confidence,
        signals=signals,
        language_candidates=language_candidates_for_text(text_for_language),
        unit_candidates=_deduplicate_unit_candidates(unit_candidates),
    )


def _block_structure_profile(document: IntermediateDocument, block: EvidenceBlock) -> StructureProfile:
    if block.kind in {"image_metadata", "embedded_visual"}:
        role: Literal["structured_table", "key_value", "narrative_note", "mixed", "visual_only", "unknown"] = "visual_only"
        confidence = 0.9
    elif block.text and block.text.strip():
        role, confidence = "narrative_note", 0.8
    else:
        role, confidence = "unknown", 0.4
    signals = [f"kind={block.kind}"]
    if block.table_id:
        signals.append(f"table_id={block.table_id}")
    if block.text:
        signals.append(f"characters={len(block.text)}")
    return StructureProfile(
        profile_id=f"profile:{block.block_id}",
        kind="block",
        role=role,
        source=block.source,
        confidence=confidence,
        signals=signals,
        language_candidates=language_candidates_for_text(block.text),
        unit_candidates=unit_candidates_for_text(block.text, block.source),
    )


def _visual_structure_profile(document: IntermediateDocument, visual: VisualAsset) -> StructureProfile:
    return StructureProfile(
        profile_id=f"profile:{visual.visual_id}",
        kind="visual",
        role="visual_only",
        source=visual.source,
        confidence=0.95,
        signals=[f"kind={visual.kind}", f"delivery_status={visual.delivery_status}"],
        language_candidates=[],
        unit_candidates=[],
    )


def _looks_like_embedded_table(text: str) -> bool:
    """A generic structural hint for a long-text fragment, not semantic OCR."""

    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if len(lines) < 3:
        return False
    delimited = sum(line.count("|") >= 2 or line.count("\t") >= 2 or line.count(",") >= 3 for line in lines)
    numeric_rows = sum(len(numeric_candidates_for_text(line)) >= 2 for line in lines)
    return delimited >= 2 or numeric_rows >= 2


def _derived_text_blocks(document: IntermediateDocument) -> list[EvidenceBlock]:
    """Add exact, source-linked long-text fragments without mutating raw evidence."""

    derived: list[EvidenceBlock] = []

    def add_fragments(
        *,
        text: str,
        source: SourcePointer,
        base_id: str,
        metadata: dict[str, Any],
    ) -> None:
        if len(text) <= MAX_DERIVED_TEXT_CHUNK_CHARACTERS:
            return
        chunk_count = math.ceil(len(text) / MAX_DERIVED_TEXT_CHUNK_CHARACTERS)
        for chunk_index in range(chunk_count):
            start = chunk_index * MAX_DERIVED_TEXT_CHUNK_CHARACTERS
            end = min(len(text), start + MAX_DERIVED_TEXT_CHUNK_CHARACTERS)
            fragment_source = source.model_copy(update={"text_span": (start, end)})
            derived.append(
                evidence_block(
                    block_id=f"derived:{base_id}:{chunk_index + 1}",
                    kind="derived_text",
                    source=fragment_source,
                    text=text[start:end],
                    metadata={
                        **metadata,
                        "chunk_index": chunk_index + 1,
                        "chunk_count": chunk_count,
                        "derived_role": "table_candidate" if _looks_like_embedded_table(text[start:end]) else "paragraph",
                    },
                )
            )

    for block in document.blocks:
        if block.kind == "derived_text" or not block.text:
            continue
        add_fragments(
            text=block.text,
            source=block.source,
            base_id=block.block_id,
            metadata={"derived_from_block_id": block.block_id},
        )
    for table in document.tables:
        for row in table.rows:
            for cell in row:
                text = cell.display_text
                if text is None or len(text) <= MAX_DERIVED_TEXT_CHUNK_CHARACTERS:
                    continue
                source = SourcePointer(
                    source_type=document.source_type,
                    file_name=document.file_name,
                    sheet_name=table.sheet_name,
                    page_number=table.page_number,
                    cell=cell.coordinate,
                    row_index=cell.row_index,
                    column_index=cell.column_index,
                    section_id=table.section_id or table.table_id,
                    parser=table.parser or document.parser,
                    extraction_confidence=table.extraction_confidence,
                )
                add_fragments(
                    text=text,
                    source=source,
                    base_id=f"{table.table_id}:{cell.coordinate}",
                    metadata={"derived_from_table_id": table.table_id, "derived_from_cell": cell.coordinate},
                )
    return derived


def _profile_line(profile: StructureProfile) -> str:
    languages = ",".join(f"{candidate.code}/{candidate.script}:{candidate.confidence}" for candidate in profile.language_candidates) or "none"
    units = ",".join(
        f"{candidate.kind}:{candidate.raw_text!r}:{candidate.currency_code or candidate.scale_multiplier}"
        for candidate in profile.unit_candidates
    ) or "none"
    return (
        f"[PROFILE id={profile.profile_id} kind={profile.kind} role={profile.role} confidence={profile.confidence} "
        f"source={_source_ref(profile.source)} languages={languages} units={units} signals={profile.signals}]"
    )


def _pack_context_lines(
    *,
    prefix: str,
    kind: Literal["document_index", "evidence"],
    lines: Iterable[tuple[str, str]],
    chunk_headers: Iterable[tuple[str, str]] = (),
) -> list[ModelContextChunk]:
    """Pack every provided line into bounded chunks, preserving line order."""

    chunks: list[ModelContextChunk] = []
    header_items = list(chunk_headers)
    current_lines: list[str] = []
    current_refs: list[str] = []
    current_size = 0
    content_lines = 0

    def seed_headers() -> None:
        nonlocal current_size
        if current_lines or not header_items:
            return
        for header_line, header_ref in header_items:
            remaining = MAX_REPEATED_CHUNK_HEADER_CHARACTERS - current_size
            if remaining <= 0:
                break
            rendered_header = header_line
            if len(rendered_header) > remaining:
                marker = "...[HEADER_TRUNCATED; FULL_METADATA_IN_CANONICAL_EVIDENCE]"
                rendered_header = rendered_header[: max(0, remaining - len(marker))] + marker
            current_lines.append(rendered_header)
            current_refs.append(header_ref)
            current_size += (1 if current_size else 0) + len(rendered_header)

    def flush() -> None:
        nonlocal current_lines, current_refs, current_size, content_lines
        if not current_lines or content_lines == 0:
            return
        text = "\n".join(current_lines)
        chunks.append(
            ModelContextChunk(
                chunk_id=f"{prefix}:{len(chunks) + 1}",
                kind=kind,
                text=text,
                character_count=len(text),
                source_refs=list(dict.fromkeys(current_refs)),
            )
        )
        current_lines, current_refs, current_size, content_lines = [], [], 0, 0

    for line, source_ref in lines:
        # Atomic evidence is normally short because long text is emitted as
        # derived fragments.  Keep a defensive path for a pathological long
        # formula or metadata value without silently deleting it.
        pieces = [line[index : index + MAX_MODEL_CONTEXT_CHUNK_CHARACTERS] for index in range(0, max(1, len(line)), MAX_MODEL_CONTEXT_CHUNK_CHARACTERS)]
        for piece in pieces:
            seed_headers()
            proposed = current_size + (1 if current_lines else 0) + len(piece)
            if content_lines and proposed > MAX_MODEL_CONTEXT_CHUNK_CHARACTERS:
                flush()
                seed_headers()
            current_lines.append(piece)
            current_refs.append(source_ref)
            current_size += (1 if current_size else 0) + len(piece)
            content_lines += 1
            if current_size >= MAX_MODEL_CONTEXT_CHUNK_CHARACTERS:
                flush()
    flush()
    return chunks


def _build_model_context_chunks(document: IntermediateDocument) -> list[ModelContextChunk]:
    document_line = (
        f"[DOCUMENT id={document.document_id or 'pending'} type={document.source_type} "
        f"file={document.file_name or 'text-input'} parser={document.parser}]"
    )
    index_lines: list[tuple[str, str]] = [(document_line, f"document:{document.document_id or 'pending'}")]
    if document.tables:
        sheet_names = list(
            dict.fromkeys(
                str(table.sheet_name)
                for table in document.tables
                if table.sheet_name
            )
        )
        index_lines.append(
            (
                "[STRUCTURE_OVERVIEW 工作簿结构 工作表清单 表格总览 "
                f"sheet_count={len(sheet_names)} sheets={sheet_names} "
                f"table_count={len(document.tables)}]",
                f"document:{document.document_id or 'pending'}",
            )
        )
        for table in document.tables:
            headers = [str(header) for header in table.headers[:20] if str(header).strip()]
            index_lines.append(
                (
                    "[TABLE_OVERVIEW 工作表 表格 数据范围 "
                    f"id={table.table_id} sheet={table.sheet_name!r} title={table.title!r} "
                    f"range={table.range!r} row_count={len(table.rows)} "
                    f"header_row={table.header_row_index} headers={headers}]",
                    table.table_id,
                )
            )
    index_lines.extend((_profile_line(profile), profile.profile_id) for profile in document.structure_profiles)
    chunks = _pack_context_lines(prefix="document-index", kind="document_index", lines=index_lines)

    derived_block_ids = {
        str(block.metadata.get("derived_from_block_id"))
        for block in document.blocks
        if block.kind == "derived_text" and block.metadata.get("derived_from_block_id")
    }
    derived_cells = {
        (str(block.metadata.get("derived_from_table_id")), str(block.metadata.get("derived_from_cell")))
        for block in document.blocks
        if block.kind == "derived_text" and block.metadata.get("derived_from_table_id") and block.metadata.get("derived_from_cell")
    }
    evidence_lines: list[tuple[str, str]] = []
    for block in document.blocks:
        if block.kind in {"worksheet", "table"}:
            continue
        if block.block_id in derived_block_ids:
            evidence_lines.append(
                (f"[BLOCK id={block.block_id} kind={block.kind} source={_source_ref(block.source)} text=LONG_TEXT_IN_DERIVED_CHUNKS]", block.block_id)
            )
            continue
        if block.text:
            evidence_lines.append(
                (f"[BLOCK id={block.block_id} kind={block.kind} source={_source_ref(block.source)}] {block.text}", block.block_id)
            )
        else:
            evidence_lines.append(
                (f"[BLOCK id={block.block_id} kind={block.kind} source={_source_ref(block.source)} metadata={block.metadata}]", block.block_id)
            )
    for visual in document.visual_assets:
        evidence_lines.append(
            (
                f"[VISUAL id={visual.visual_id} kind={visual.kind} source={_source_ref(visual.source)} "
                f"status={visual.delivery_status} metadata={_searchable_visual_metadata(visual.metadata)}]",
                visual.visual_id,
            )
        )
    chunks.extend(_pack_context_lines(prefix="evidence", kind="evidence", lines=evidence_lines))
    for table_index, table in enumerate(document.tables, start=1):
        table_ref = table.table_id
        table_origin_column = min((cell.column_index for row in table.rows for cell in row), default=1)
        metadata = [
            f"id={table.table_id}",
            f"sheet={table.sheet_name}" if table.sheet_name else None,
            f"page={table.page_number}" if table.page_number else None,
            f"state={table.sheet_state}" if table.sheet_state else None,
            f"range={table.range}" if table.range else None,
            f"parser={table.parser}" if table.parser else None,
        ]
        table_headers: list[tuple[str, str]] = [
            (f"[TABLE {' '.join(item for item in metadata if item)}]", table_ref)
        ]
        if table.merged_ranges:
            merge_sample = table.merged_ranges[:50]
            omitted = len(table.merged_ranges) - len(merge_sample)
            suffix = f" | omitted={omitted}; full_list=canonical_evidence" if omitted else ""
            table_headers.append(
                (f"[MERGED_RANGES count={len(table.merged_ranges)}] {', '.join(merge_sample)}{suffix}", table_ref)
            )
        column_schema = [
            f"{get_column_letter(table_origin_column + index)}={header}"
            for index, header in enumerate(table.headers)
            if header
        ]
        if column_schema:
            table_headers.append((f"[COLUMNS] {' | '.join(column_schema)}", table_ref))
        table_lines: list[tuple[str, str]] = []
        for row in table.rows:
            rendered_cells: list[str] = []
            for cell in row:
                if cell.display_text is None and cell.formula is None:
                    continue
                if (table.table_id, cell.coordinate) in derived_cells:
                    rendered_value = "LONG_TEXT_IN_DERIVED_CHUNKS"
                else:
                    rendered_value = repr(cell.display_text)
                formula = f" formula={cell.formula!r}" if cell.formula else ""
                rendered_cells.append(f"{cell.coordinate}={rendered_value}{formula}")
            if rendered_cells:
                row_index = min(cell.row_index for cell in row)
                row_source = SourcePointer(
                    source_type=document.source_type,
                    file_name=document.file_name,
                    sheet_name=table.sheet_name,
                    page_number=table.page_number,
                    row_index=row_index,
                    section_id=table.section_id or table.table_id,
                    parser=table.parser or document.parser,
                    extraction_confidence=table.extraction_confidence,
                )
                table_lines.append(
                    (
                        f"[ROW source={_source_ref(row_source)}] {' | '.join(rendered_cells)}",
                        f"{table.table_id}:row:{row_index}",
                    )
                )
        chunks.extend(
            _pack_context_lines(
                prefix=f"table-{table_index}",
                kind="evidence",
                lines=table_lines,
                chunk_headers=table_headers,
            )
        )
    return chunks


def _enrich_intermediate_document(document: IntermediateDocument) -> None:
    """Attach generic evidence descriptors without making business mappings."""

    derived = _derived_text_blocks(document)
    if derived:
        document.blocks.extend(derived)
    profiles: list[StructureProfile] = [_table_structure_profile(document, table) for table in document.tables]
    profiles.extend(_block_structure_profile(document, block) for block in document.blocks if block.kind != "derived_text")
    profiles.extend(_visual_structure_profile(document, visual) for visual in document.visual_assets)
    document_text = "\n".join(
        [document.raw_text or "", *[" ".join((cell.display_text or "") for row in table.rows[:3] for cell in row[:12]) for table in document.tables[:12]]]
    )
    document_role: Literal["structured_table", "key_value", "narrative_note", "mixed", "visual_only", "unknown"]
    if document.tables and (document.blocks or document.visual_assets):
        document_role = "mixed"
    elif document.tables:
        document_role = "structured_table"
    elif document.blocks:
        document_role = "narrative_note"
    elif document.visual_assets:
        document_role = "visual_only"
    else:
        document_role = "unknown"
    profiles.insert(
        0,
        StructureProfile(
            profile_id="profile:document",
            kind="document",
            role=document_role,
            source=SourcePointer(source_type=document.source_type, file_name=document.file_name, parser=document.parser),
            confidence=0.8,
            signals=[f"tables={len(document.tables)}", f"blocks={len(document.blocks)}", f"visual_assets={len(document.visual_assets)}"],
            language_candidates=language_candidates_for_text(document_text),
            unit_candidates=[],
        ),
    )
    document.structure_profiles = profiles
    document.model_context_chunks = _build_model_context_chunks(document)


def render_intermediate_for_model(
    document: IntermediateDocument,
    *,
    max_characters: int = MAX_MODEL_CONTEXT_CHARACTERS,
    chunk_ids: Iterable[str] | None = None,
) -> str:
    """Render evidence structurally for an LLM without replacing the evidence.

    Tables retain row and cell coordinates.  The document keeps an *unbounded*
    ordered ``model_context_chunks`` collection; this convenience rendering is
    intentionally bounded and explicitly says when more chunks are available.
    """

    limit = max(1_000, max_characters)
    if document.model_context_chunks:
        requested = set(chunk_ids) if chunk_ids is not None else None
        chunks = [chunk for chunk in document.model_context_chunks if requested is None or chunk.chunk_id in requested]
        lines: list[str] = []
        for index, chunk in enumerate(chunks):
            candidate = "\n".join([*lines, chunk.text])
            if len(candidate) > limit:
                remaining = len(chunks) - index
                marker = (
                    f"[MODEL_CONTEXT_PARTIAL: {remaining} source-linked chunks remain; "
                    "select their model_context_chunks by chunk_id. Canonical Evidence JSON is complete.]"
                )
                if len("\n".join([*lines, marker])) <= limit:
                    lines.append(marker)
                return "\n".join(lines)
            lines.append(chunk.text)
        return "\n".join(lines)

    # Pre-finalisation fallback for callers constructing a document directly.
    lines = [
        f"[DOCUMENT id={document.document_id or 'pending'} type={document.source_type} file={document.file_name or 'text-input'} parser={document.parser}]"
    ]

    def append(line: str) -> bool:
        candidate = "\n".join([*lines, line])
        if len(candidate) > limit:
            lines.append("[MODEL_CONTEXT_TRUNCATED: canonical Evidence JSON remains complete]")
            return False
        lines.append(line)
        return True

    for block in document.blocks:
        if block.kind in {"worksheet", "table"}:
            continue
        location = block.source.section_id or (
            f"page:{block.source.page_number}" if block.source.page_number else "source"
        )
        text = (block.text or "").strip().replace("\r\n", "\n")
        if text:
            if not append(f"[BLOCK id={block.block_id} kind={block.kind} location={location}] {text}"):
                return "\n".join(lines)
        else:
            if not append(f"[BLOCK id={block.block_id} kind={block.kind} location={location} metadata={block.metadata}]"):
                return "\n".join(lines)

    for visual in document.visual_assets:
        location = visual.source.cell or visual.source.sheet_name or "source"
        if not append(
            f"[VISUAL id={visual.visual_id} kind={visual.kind} location={location} "
            f"status={visual.delivery_status} metadata={_searchable_visual_metadata(visual.metadata)}]"
        ):
            return "\n".join(lines)

    for table in document.tables:
        metadata = [
            f"id={table.table_id}",
            f"sheet={table.sheet_name}" if table.sheet_name else None,
            f"page={table.page_number}" if table.page_number else None,
            f"state={table.sheet_state}" if table.sheet_state else None,
            f"range={table.range}" if table.range else None,
            f"parser={table.parser}" if table.parser else None,
        ]
        if not append(f"[TABLE {' '.join(item for item in metadata if item)}]"):
            return "\n".join(lines)
        if table.merged_ranges and not append(f"[MERGED_RANGES] {', '.join(table.merged_ranges)}"):
            return "\n".join(lines)
        for row in table.rows:
            values = [
                f"{cell.coordinate}={cell.display_text!r}" + (f" formula={cell.formula!r}" if cell.formula else "")
                for cell in row
                if cell.display_text is not None or cell.formula is not None
            ]
            if values and not append("[ROW] " + " | ".join(values)):
                return "\n".join(lines)
    return "\n".join(lines)


def finalize_intermediate_evidence(document: IntermediateDocument, content: bytes) -> IntermediateDocument:
    """Attach deterministic request-only identity and bounded model rendering."""

    digest = hashlib.sha256(content).hexdigest()
    document.document_id = f"doc-{digest[:16]}"
    document.file_sha256 = digest
    document.file_size_bytes = len(content)
    _enrich_intermediate_document(document)
    document.model_context = render_intermediate_for_model(document)
    return document


def _metric_for_label(value: Any) -> MetricDefinition | None:
    if not isinstance(value, str):
        return None
    normalised = _normalised_label(value)
    if not normalised:
        return None
    direct = _METRIC_BY_NORMALISED_ALIAS.get(normalised)
    if direct:
        return direct
    # Deliberately only accept an unambiguous contained alias.  A future model
    # can propose mappings for the remaining labels, but it must not silently
    # turn arbitrary fields into accounting facts.
    candidates = {
        metric.code: metric
        for alias, metric in _METRIC_BY_NORMALISED_ALIAS.items()
        if len(alias) >= 4 and alias in normalised
    }
    return next(iter(candidates.values())) if len(candidates) == 1 else None


def _semantic_mapping_index(
    mappings: Iterable[SemanticLabelMapping],
    table_id: str,
) -> dict[str, SemanticLabelMapping]:
    """Keep the strongest local-model candidate for each uploaded label."""

    indexed: dict[str, SemanticLabelMapping] = {}
    for mapping in mappings:
        if mapping.table_id is not None and mapping.table_id != table_id:
            continue
        label = _normalised_label(mapping.label)
        if not label or (mapping.target not in _METRIC_BY_CODE and mapping.target not in _STRUCTURAL_TARGETS):
            continue
        previous = indexed.get(label)
        if previous is None or mapping.confidence > previous.confidence:
            indexed[label] = mapping
    return indexed


def _metric_resolution(
    value: Any,
    semantic_index: dict[str, SemanticLabelMapping],
) -> tuple[MetricDefinition, Literal["rule", "semantic_candidate"], float] | None:
    metric = _metric_for_label(value)
    if metric is not None:
        return metric, "rule", 1.0
    semantic = semantic_index.get(_normalised_label(str(value or "")))
    if semantic is None:
        return None
    metric = _METRIC_BY_CODE.get(semantic.target)
    if metric is None:
        return None
    return metric, "semantic_candidate", semantic.confidence


def _header_role(value: Any, semantic_index: dict[str, SemanticLabelMapping]) -> str | None:
    label = _normalised_label(str(value or ""))
    if label in {_normalised_label(item) for item in _PERIOD_HEADER_LABELS}:
        return "__PERIOD__"
    if label in {_normalised_label(item) for item in _ITEM_HEADER_LABELS}:
        return "__ITEM__"
    if any(_normalised_label(item) in label for item in _PRIOR_AMOUNT_HEADER_LABELS):
        return "__PRIOR_AMOUNT__"
    if any(_normalised_label(item) in label for item in _CURRENT_AMOUNT_HEADER_LABELS):
        return "__CURRENT_AMOUNT__"
    semantic = semantic_index.get(label)
    return semantic.target if semantic and semantic.target in _STRUCTURAL_TARGETS else None


def _period_from_value(value: Any, default_year: str | None = None) -> str | None:
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m")
    # PDF table extractors commonly insert spaces between Chinese date tokens
    # (for example ``2024 年 12 月 31 日``).  Those spaces carry no period
    # meaning and should not prevent structural date recognition.
    text = re.sub(r"\s+", "", str(value or "").strip())
    if re.fullmatch(r"\d{4}-\d{2}", text):
        month = int(text[-2:])
        return text if 1 <= month <= 12 else None
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
        try:
            return date.fromisoformat(text).strftime("%Y-%m")
        except ValueError:
            return None
    if re.fullmatch(r"\d{4}", text):
        return text
    match = re.search(r"(?P<year>20\d{2})年(?:(?P<month>1[0-2]|0?[1-9])月)?", text)
    if match:
        return f"{match.group('year')}-{int(match.group('month')):02d}" if match.group("month") else match.group("year")
    month_match = re.fullmatch(r"(?P<month>1[0-2]|0?[1-9])月", text)
    if month_match and default_year and re.fullmatch(r"20\d{2}", default_year):
        return f"{default_year}-{int(month_match.group('month')):02d}"
    return None


def _previous_period(period: str | None) -> str | None:
    if period is None:
        return None
    if re.fullmatch(r"\d{4}", period):
        return str(int(period) - 1)
    if re.fullmatch(r"\d{4}-\d{2}", period):
        year, month = (int(part) for part in period.split("-"))
        return f"{year - 1}-12" if month == 1 else f"{year}-{month - 1:02d}"
    return None


def _number(value: Any) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        candidate = float(value)
        return candidate if math.isfinite(candidate) else None
    text = str(value).strip().replace(",", "")
    if text.startswith("(") and text.endswith(")"):
        text = f"-{text[1:-1]}"
    try:
        candidate = float(text)
    except ValueError:
        return None
    return candidate if math.isfinite(candidate) else None


def _issue(
    issues: list[ValidationIssue],
    severity: Literal["error", "warning"],
    code: str,
    message: str,
    source: SourcePointer | None = None,
) -> None:
    issues.append(ValidationIssue(severity=severity, code=code, message=message, source=source))


def _validate_facts(facts: Iterable[StandardFinancialFact], issues: list[ValidationIssue]) -> None:
    seen: set[tuple[str, str | None]] = set()
    for fact in facts:
        identity = (fact.standard_item_code, fact.period)
        if identity in seen:
            _issue(issues, "warning", "duplicate_fact", f"检测到重复指标：{fact.standard_item_name} / {fact.period or '期间未知'}。", fact.source)
        seen.add(identity)
        if fact.value is None:
            _issue(issues, "error", "invalid_amount", f"{fact.original_item_name} 未能解析为有限数值。", fact.source)
        if not fact.period:
            _issue(issues, "warning", "period_missing", f"{fact.original_item_name} 缺少报告期间，需客户确认。", fact.source)


def _build_fact(
    metric: MetricDefinition,
    original_item_name: str,
    value: Any,
    period: str | None,
    source: SourcePointer,
    *,
    confidence: float = 1.0,
    mapping_method: Literal["rule", "semantic_candidate", "customer_confirmed"] = "rule",
    unit: Literal["yuan", "unknown"] = "yuan",
) -> StandardFinancialFact:
    return StandardFinancialFact(
        statement_type=metric.statement_type,
        standard_item_code=metric.code,
        standard_item_name=metric.name,
        original_item_name=original_item_name,
        period=period,
        value=_number(value),
        source=source,
        mapping_method=mapping_method,
        mapping_confidence=confidence,
        unit=unit,
    )


def _make_excel_cell(
    *,
    row_index: int,
    column_index: int,
    formula_value: Any,
    cached_value: Any,
) -> IntermediateCell:
    def formula_expression(value: Any) -> str | None:
        if isinstance(value, str) and value.startswith("="):
            return value
        text = getattr(value, "text", None)
        if isinstance(text, str) and text.startswith("="):
            return text
        # Excel What-If Data Tables are stored by OOXML as a formula object
        # with inputs and a range, not a normal A1 formula string.  ``str``
        # embeds a Python memory address, which would make Evidence JSON
        # nondeterministic.  Preserve a stable, explicit descriptor instead.
        if value.__class__.__name__ == "DataTableFormula":
            fields = [
                ("ref", getattr(value, "ref", None)),
                ("row_input", getattr(value, "r1", None)),
                ("column_input", getattr(value, "r2", None)),
                ("two_variable", getattr(value, "dt2D", None)),
            ]
            details = ",".join(f"{name}={item}" for name, item in fields if item not in {None, ""})
            return f"=DATA_TABLE({details})"
        return None

    # ``openpyxl`` represents normal formulas as strings, but dynamic/array
    # formulas as formula objects (for example ``ArrayFormula`` with a
    # ``text`` attribute).  Treat both as formulas.  Otherwise an array
    # formula becomes an opaque object string in the Evidence JSON and loses
    # its executable expression.
    formula = formula_expression(formula_value)
    display_text = _display(cached_value if cached_value is not None else formula_value)
    return IntermediateCell(
        coordinate=f"{get_column_letter(column_index)}{row_index}",
        row_index=row_index,
        column_index=column_index,
        value=None if formula else _serialise_cell_value(formula_value),
        display_text=display_text,
        formula=formula,
        cached_value=_serialise_cell_value(cached_value),
        numeric_candidates=numeric_candidates_for_text(display_text),
    )


def _is_nonempty_row(cells: list[IntermediateCell]) -> bool:
    return any(cell.value is not None or cell.formula is not None or cell.cached_value is not None for cell in cells)


def _headers_for_rows(rows: list[list[IntermediateCell]]) -> list[str]:
    generic_candidates: list[tuple[int, int, list[str]]] = []
    origin_column = min((cell.column_index for row in rows for cell in row), default=1)
    for row in rows[:30]:
        max_column = max((cell.column_index for cell in row), default=0)
        values = [""] * max(0, max_column - origin_column + 1)
        for cell in row:
            values[cell.column_index - origin_column] = cell.display_text or ""
        metrics = sum(_metric_for_label(value) is not None for value in values)
        has_period = any(_normalised_label(value) in {"period", "月份", "期间", "报告期间", "日期", "date", "month"} for value in values)
        has_item = any(_normalised_label(value) in {"项目", "科目", "item"} for value in values)
        if metrics >= 2 or has_period or (has_item and metrics >= 1):
            return values
        nonempty = [value for value in values if value.strip()]
        textual = [
            value
            for cell, value in zip(row, values)
            if value.strip() and isinstance(cell.value, str)
        ]
        if len(nonempty) >= 2 and len(textual) / len(nonempty) >= 0.6 and len(set(nonempty)) == len(nonempty):
            generic_candidates.append((len(textual), len(nonempty), values))
    if generic_candidates:
        return max(generic_candidates, key=lambda item: (item[0], item[1]))[2]
    return []


_STRUCTURAL_PERIOD_HEADERS = {"period", "月份", "期间", "报告期间", "日期", "date", "month"}
_STRUCTURAL_ITEM_HEADERS = {"项目", "科目", "item", "项目名称", "科目名称"}
_STRUCTURAL_AMOUNT_HEADERS = {
    "本期金额",
    "本年累计数",
    "本期数",
    "期末余额",
    "amount",
    "金额",
    "发生额",
    "上期金额",
    "上年累计数",
    "上期数",
    "期初余额",
    "prioramount",
    "上年金额",
}


def _cell_has_content(cell: IntermediateCell) -> bool:
    """Treat empty strings from legacy XLS as layout whitespace, not data."""

    if cell.formula:
        return True
    for value in (cell.display_text, cell.value, cell.cached_value):
        if value is not None and str(value).strip():
            return True
    return False


def _contiguous_ranges(indices: list[int]) -> list[tuple[int, int]]:
    if not indices:
        return []
    ranges: list[tuple[int, int]] = []
    start = previous = indices[0]
    for current in indices[1:]:
        if current != previous + 1:
            ranges.append((start, previous))
            start = current
        previous = current
    ranges.append((start, previous))
    return ranges


def _is_structural_header_row(row: list[IntermediateCell]) -> bool:
    labels = [_normalised_label((cell.display_text or "").strip()) for cell in row]
    nonempty = [label for label in labels if label]
    if len(nonempty) < 2:
        return False
    has_period = any(label in _STRUCTURAL_PERIOD_HEADERS for label in nonempty)
    has_item = any(label in _STRUCTURAL_ITEM_HEADERS for label in nonempty)
    has_amount = any(label in _STRUCTURAL_AMOUNT_HEADERS for label in nonempty)
    return has_period or (has_item and has_amount)


def _first_structural_header(rows: list[list[IntermediateCell]]) -> tuple[int | None, list[IntermediateCell] | None]:
    for index, row in enumerate(rows):
        if _is_structural_header_row(row):
            return index, row
    return None, None


def _title_cell_before_header(rows: list[list[IntermediateCell]], header_index: int | None) -> IntermediateCell | None:
    if header_index is None:
        return None
    for row in reversed(rows[:header_index]):
        populated = [cell for cell in row if _cell_has_content(cell)]
        if not populated:
            continue
        return populated[0] if len(populated) == 1 else None
    return None


def _merged_ranges_within(
    merged_ranges: list[str],
    *,
    row_low: int,
    row_high: int,
    column_low: int,
    column_high: int,
) -> list[str]:
    selected: list[str] = []
    for item in merged_ranges:
        match = re.fullmatch(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", item)
        if match is None:
            continue
        start_column, start_row = column_index_from_string(match.group(1)), int(match.group(2))
        end_column, end_row = column_index_from_string(match.group(3)), int(match.group(4))
        if row_low <= start_row and end_row <= row_high and column_low <= start_column and end_column <= column_high:
            selected.append(item)
    return selected


def _row_is_title_like(row: list[IntermediateCell]) -> bool:
    """Whether a row is a single-label title rather than table content."""

    return len([cell for cell in row if _cell_has_content(cell)]) == 1


def _starts_new_table(
    rows: list[list[IntermediateCell]],
    header_index: int,
) -> bool:
    """Recognise a restarted table without splitting every visual spacer.

    A blank row on its own is common in forms, explanatory sheets and styled
    models.  It is not proof of a new table.  A later recognised header must
    be separated by layout whitespace (optionally followed by a one-cell
    title) before it starts an independent logical table.
    """

    if header_index == 0:
        return True
    current_row = rows[header_index][0].row_index
    previous_row = rows[header_index - 1][0].row_index
    if current_row - previous_row > 1:
        return True
    if header_index >= 2:
        prior_row = rows[header_index - 2][0].row_index
        if previous_row - prior_row > 1 and _row_is_title_like(rows[header_index - 1]):
            return True
    return False


def _horizontal_table_ranges(rows: list[list[IntermediateCell]]) -> list[tuple[int, int]]:
    """Return side-by-side table ranges only when every range has a header.

    This intentionally avoids treating decorative whitespace and notes in a
    complex financial model as hundreds of unrelated tables.
    """

    active_columns = sorted(
        {
            cell.column_index
            for row in rows
            for cell in row
            if _cell_has_content(cell)
        }
    )
    column_ranges = _contiguous_ranges(active_columns)
    if len(column_ranges) < 2:
        return column_ranges

    qualifying_ranges: list[tuple[int, int]] = []
    for column_low, column_high in column_ranges:
        local_rows = [
            [cell for cell in row if column_low <= cell.column_index <= column_high]
            for row in rows
        ]
        if any(_is_structural_header_row(row) for row in local_rows):
            qualifying_ranges.append((column_low, column_high))
    return column_ranges if len(qualifying_ranges) == len(column_ranges) else [(active_columns[0], active_columns[-1])]


def split_sheet_into_tables(
    table: IntermediateTable,
    *,
    source_type: Literal["excel"],
    file_name: str,
) -> list[IntermediateTable]:
    """Split one worksheet grid into independent, disconnected table regions.

    A blank row/column is merely layout whitespace in many real workbooks, so
    it cannot by itself define a table.  A section is split only when a new
    recognised table header restarts after layout whitespace.  Side-by-side
    tables are split only when each side independently contains a recognised
    header.  This preserves genuine stacked reports while avoiding a cascade
    of one-cell fragments from notes and financial-model layouts.
    """

    source_rows = table.rows
    source_rows = [row for row in source_rows if row and any(_cell_has_content(cell) for cell in row)]
    if not source_rows:
        return [table]

    all_structural_headers = [index for index, row in enumerate(source_rows) if _is_structural_header_row(row)]
    if not all_structural_headers:
        return [table]
    structural_headers = [all_structural_headers[0]] + [
        index for index in all_structural_headers[1:] if _starts_new_table(source_rows, index)
    ]

    # Preserve all sheet evidence.  The first section includes any explanatory
    # material before its first recognised header; later sections start at an
    # immediately preceding title where present.
    vertical_starts: list[int] = [0]
    for header_index in structural_headers[1:]:
        start_index = header_index
        if _row_is_title_like(source_rows[header_index - 1]):
            start_index -= 1
        if start_index > vertical_starts[-1]:
            vertical_starts.append(start_index)

    sections: list[IntermediateTable] = []
    parent_table_id = table.parent_table_id or table.table_id
    for vertical_index, start_index in enumerate(vertical_starts):
        end_index = vertical_starts[vertical_index + 1] - 1 if vertical_index + 1 < len(vertical_starts) else len(source_rows) - 1
        rows_in_band = source_rows[start_index : end_index + 1]
        row_low, row_high = rows_in_band[0][0].row_index, rows_in_band[-1][0].row_index
        for column_low, column_high in _horizontal_table_ranges(rows_in_band):
            section_rows = [
                [cell for cell in row if column_low <= cell.column_index <= column_high]
                for row in rows_in_band
            ]
            if not any(_cell_has_content(cell) for row in section_rows for cell in row):
                continue
            section_index = len(sections) + 1
            section_id = f"{parent_table_id}:section-{section_index}"
            header_index, header = _first_structural_header(section_rows)
            title_cell = _title_cell_before_header(section_rows, header_index)
            title_source = (
                SourcePointer(
                    source_type=source_type,
                    file_name=file_name,
                    sheet_name=table.sheet_name,
                    cell=title_cell.coordinate,
                    row_index=title_cell.row_index,
                    column_index=title_cell.column_index,
                    section_id=section_id,
                    parser=table.parser,
                    extraction_confidence=table.extraction_confidence,
                )
                if title_cell is not None
                else None
            )
            title = (title_cell.display_text or "").strip() if title_cell else table.title
            explicit_historical = bool(title and "含往年" in re.sub(r"\s+", "", title))
            scope: Literal["includes_prior_periods", "current_period_only", "unknown"] = "includes_prior_periods" if explicit_historical else "unknown"
            sections.append(
                IntermediateTable(
                    table_id=f"{parent_table_id}:table-{section_index}",
                    parent_table_id=parent_table_id,
                    section_id=section_id,
                    title=title,
                    title_source=title_source,
                    sheet_name=table.sheet_name,
                    page_number=table.page_number,
                    sheet_state=table.sheet_state,
                    range=f"{get_column_letter(column_low)}{row_low}:{get_column_letter(column_high)}{row_high}",
                    parser=table.parser,
                    extraction_confidence=table.extraction_confidence,
                    header_row_index=header[0].row_index if header else None,
                    header_range=(
                        f"{get_column_letter(column_low)}{header[0].row_index}:{get_column_letter(column_high)}{header[0].row_index}"
                        if header
                        else None
                    ),
                    scope=scope,
                    scope_confidence=1.0 if explicit_historical else 0.0,
                    scope_evidence=title_source if explicit_historical else None,
                    headers=_headers_for_rows([header]) if header else _headers_for_rows(section_rows),
                    rows=section_rows,
                    merged_ranges=_merged_ranges_within(
                        table.merged_ranges,
                        row_low=row_low,
                        row_high=row_high,
                        column_low=column_low,
                        column_high=column_high,
                    ),
                )
            )

    # A sibling whose otherwise identical title only removes “含往年” is a
    # strong layout-level current-period candidate.  Keep the evidence and a
    # confidence rather than asserting this for unrelated sheets.
    historical_titles = {
        re.sub(r"含往年|\s+|[（）()]", "", item.title or "")
        for item in sections
        if item.scope == "includes_prior_periods" and item.title
    }
    classified: list[IntermediateTable] = []
    for item in sections:
        signature = re.sub(r"含往年|\s+|[（）()]", "", item.title or "")
        if item.scope == "unknown" and signature and signature in historical_titles:
            classified.append(
                item.model_copy(
                    update={
                        "scope": "current_period_only",
                        "scope_confidence": 0.86,
                        "scope_evidence": item.title_source,
                    }
                )
            )
        else:
            classified.append(item)
    return classified or [table]


def _validate_tabular_upload(file_name: str, content: bytes) -> str:
    suffix = Path(file_name).suffix.lower()
    if suffix not in {".xlsx", ".xls", ".csv"}:
        raise FinanceIntakeError("表格入口仅接收 .xlsx、.xls 或 .csv 文件。")
    if not content:
        raise FinanceIntakeError("上传的表格文件为空。")
    if len(content) > MAX_EXCEL_BYTES:
        raise FinanceIntakeError(f"表格文件超过 {MAX_EXCEL_BYTES // (1024 * 1024)} MB 限制。")
    return suffix


def _drawing_anchor(anchor: Any) -> tuple[str | None, str | None, int | None, int | None]:
    """Return a readable cell anchor/range and approximate drawing size."""

    if isinstance(anchor, str):
        return anchor, anchor, None, None
    start = getattr(anchor, "_from", None)
    if start is None:
        return None, None, None, None
    start_cell = f"{get_column_letter(start.col + 1)}{start.row + 1}"
    end = getattr(anchor, "to", None) or getattr(anchor, "_to", None)
    end_cell = f"{get_column_letter(end.col + 1)}{end.row + 1}" if end is not None else start_cell
    extent = getattr(anchor, "ext", None)
    width_px = round(extent.cx / 9_525) if extent is not None and getattr(extent, "cx", None) else None
    height_px = round(extent.cy / 9_525) if extent is not None and getattr(extent, "cy", None) else None
    return start_cell, f"{start_cell}:{end_cell}", width_px, height_px


def _reference_values(workbook: Any, reference: str | None) -> list[Any]:
    """Resolve a simple chart series reference from the cached workbook."""

    if not reference:
        return []
    try:
        sheet_name, (min_column, min_row, max_column, max_row) = range_to_tuple(reference)
        worksheet = workbook[sheet_name]
    except Exception:
        return []
    return [
        worksheet.cell(row=row, column=column).value
        for row in range(min_row, max_row + 1)
        for column in range(min_column, max_column + 1)
    ]


def _series_reference(series: Any, attribute: str) -> str | None:
    """Read a chart XML reference without guessing a cell address."""

    source = getattr(series, attribute, None)
    if source is None:
        return None
    for reference_name in ("numRef", "strRef", "multiLvlStrRef"):
        reference = getattr(source, reference_name, None)
        formula = getattr(reference, "f", None) if reference is not None else None
        if isinstance(formula, str) and formula:
            return formula
    return None


def _chart_title(chart: Any, workbook: Any) -> str | None:
    title = getattr(chart, "title", None)
    if title is None:
        return None
    try:
        text_ref = title.tx.strRef.f if title.tx and title.tx.strRef else None
        if text_ref:
            values = _reference_values(workbook, text_ref)
            if values and values[0] is not None:
                return _display(values[0])
        paragraphs = title.tx.rich.p if title.tx and title.tx.rich else []
        text = "".join(
            run.t
            for paragraph in paragraphs
            for run in (getattr(paragraph, "r", None) or [])
            if getattr(run, "t", None)
        ).strip()
        return text or None
    except Exception:
        return None


def _series_title(series: Any, workbook: Any, fallback: str) -> str:
    try:
        text_ref = series.tx.strRef.f if series.tx and series.tx.strRef else None
        if text_ref:
            values = _reference_values(workbook, text_ref)
            if values and values[0] is not None:
                return _display(values[0]) or fallback
        literal = series.tx.v if series.tx else None
        return _display(literal) or fallback
    except Exception:
        return fallback


def _chart_series(chart: Any, workbook: Any) -> list[dict[str, Any]]:
    series_data: list[dict[str, Any]] = []
    for index, series in enumerate(getattr(chart, "ser", []) or [], start=1):
        value_ref = _series_reference(series, "val")
        category_ref = _series_reference(series, "cat")
        values = _reference_values(workbook, value_ref)
        categories = _reference_values(workbook, category_ref)
        numeric_values = [_number(value) for value in values]
        if not any(value is not None for value in numeric_values):
            continue
        series_data.append(
            {
                "name": _series_title(series, workbook, f"Series {index}"),
                "value_reference": value_ref,
                "category_reference": category_ref,
                "values": numeric_values,
                "categories": [_display(value) or "" for value in categories],
            }
        )
    return series_data


def _render_chart_preview(chart: Any, workbook: Any) -> tuple[bytes | None, dict[str, Any]]:
    """Create a bounded visual preview for common native Excel chart types.

    The preview is a visual aid for Qwen3-VL, never a replacement for the
    source cell values.  Unsupported charts remain fully represented as
    metadata and are explicitly marked for a future Office renderer.
    """

    chart_type = chart.__class__.__name__
    normalized_type = chart_type.lower()
    series = _chart_series(chart, workbook)
    metadata: dict[str, Any] = {
        "chart_class": chart_type,
        "chart_title": _chart_title(chart, workbook),
        "series": [
            {
                "name": item["name"],
                "value_reference": item["value_reference"],
                "category_reference": item["category_reference"],
            }
            for item in series
        ],
        "preview_is_reconstruction": True,
    }
    supported = {"barchart", "linechart", "areachart", "scatterchart", "piechart", "doughnutchart"}
    if normalized_type not in supported:
        metadata["rendering_reason"] = "unsupported_chart_type"
        return None, metadata
    if not series:
        metadata["rendering_reason"] = "no_resolvable_numeric_series"
        return None, metadata
    try:
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt

        figure, axis = plt.subplots(figsize=(7.2, 4.2), dpi=120)
        title = metadata["chart_title"] or chart_type
        max_points = 48
        if normalized_type in {"piechart", "doughnutchart"}:
            first = series[0]
            values = [value for value in first["values"][:max_points] if value is not None and value >= 0]
            labels = (first["categories"] or [str(index + 1) for index in range(len(values))])[: len(values)]
            if not values or sum(values) <= 0:
                plt.close(figure)
                metadata["rendering_reason"] = "pie_has_no_positive_values"
                return None, metadata
            wedges, _ = axis.pie(values, labels=labels, autopct="%1.0f%%")
            if normalized_type == "doughnutchart":
                for wedge in wedges:
                    wedge.set_width(0.45)
        elif normalized_type == "scatterchart":
            for index, item in enumerate(series):
                values = item["values"][:max_points]
                categories = item["categories"][:max_points]
                x_values = [_number(value) for value in categories]
                if not any(value is not None for value in x_values):
                    x_values = list(range(1, len(values) + 1))
                axis.scatter(x_values, values, label=item["name"])
            axis.legend(loc="best")
        elif normalized_type == "barchart":
            count = max(len(item["values"][:max_points]) for item in series)
            labels = (series[0]["categories"] or [str(index + 1) for index in range(count)])[:count]
            width = 0.8 / max(1, len(series))
            positions = list(range(count))
            for index, item in enumerate(series):
                values = [(value if value is not None else 0.0) for value in item["values"][:count]]
                offsets = [position - 0.4 + width / 2 + index * width for position in positions]
                axis.barh(offsets, values, height=width, label=item["name"]) if getattr(chart, "barDir", "col") == "bar" else axis.bar(offsets, values, width=width, label=item["name"])
            if getattr(chart, "barDir", "col") == "bar":
                axis.set_yticks(positions, labels)
            else:
                axis.set_xticks(positions, labels, rotation=30, ha="right")
            axis.legend(loc="best")
        else:
            for item in series:
                values = item["values"][:max_points]
                labels = item["categories"][:max_points]
                positions = list(range(len(values)))
                if normalized_type == "areachart":
                    axis.fill_between(positions, values, alpha=0.28, label=item["name"])
                    axis.plot(positions, values)
                else:
                    axis.plot(positions, values, marker="o", label=item["name"])
                if labels:
                    axis.set_xticks(positions, labels, rotation=30, ha="right")
            axis.legend(loc="best")
        axis.set_title(title)
        axis.grid(axis="y", alpha=0.25)
        figure.tight_layout()
        output = BytesIO()
        figure.savefig(output, format="png", dpi=120)
        plt.close(figure)
        image = output.getvalue()
        metadata["rendering_reason"] = "rendered"
        metadata["preview_point_limit"] = max_points
        return image, metadata
    except Exception as exc:
        metadata["rendering_reason"] = f"render_failed:{type(exc).__name__}"
        return None, metadata


def _xlsx_visual_assets_for_sheet(
    worksheet: Any,
    cached_workbook: Any,
    *,
    file_name: str,
    asset_offset: int,
    asset_limit: int,
) -> list[VisualAsset]:
    """Extract native images and render common native charts in memory only."""

    assets: list[VisualAsset] = []
    source_file = Path(file_name).name
    for image_index, image in enumerate(getattr(worksheet, "_images", []) or [], start=1):
        if len(assets) >= asset_limit:
            break
        anchor_cell, anchor_range, anchor_width, anchor_height = _drawing_anchor(getattr(image, "anchor", None))
        try:
            raw = image._data()
            if not raw or len(raw) > MAX_EXCEL_VISUAL_ASSET_BYTES:
                raise ValueError("image_bytes_out_of_bounds")
            from PIL import Image

            with Image.open(BytesIO(raw)) as opened:
                width_px, height_px = opened.size
                image_format = (opened.format or getattr(image, "format", "png")).lower()
            media_type = f"image/{'jpeg' if image_format in {'jpg', 'jpeg'} else image_format}"
            status: Literal["ready_for_vision", "metadata_only", "rendering_unavailable"] = "ready_for_vision"
            image_bytes: bytes | None = raw
            reason = None
        except Exception as exc:
            width_px, height_px = anchor_width, anchor_height
            media_type = None
            status = "metadata_only"
            image_bytes = None
            image_format = None
            reason = type(exc).__name__
        visual_id = f"excel:{worksheet.title}:image-{asset_offset + image_index}"
        assets.append(
            VisualAsset(
                visual_id=visual_id,
                kind="embedded_image",
                source=SourcePointer(
                    source_type="excel",
                    file_name=source_file,
                    sheet_name=worksheet.title,
                    cell=anchor_cell,
                    parser="openpyxl-drawing",
                    extraction_confidence=1.0,
                ),
                media_type=media_type,
                width_px=width_px,
                height_px=height_px,
                sha256=hashlib.sha256(image_bytes).hexdigest() if image_bytes else None,
                delivery_status=status,
                metadata={"anchor_range": anchor_range, "image_format": image_format, "read_error": reason},
                image_bytes=image_bytes,
            )
        )
    chart_offset = asset_offset + len(getattr(worksheet, "_images", []) or [])
    for chart_index, chart in enumerate(getattr(worksheet, "_charts", []) or [], start=1):
        if len(assets) >= asset_limit:
            break
        anchor_cell, anchor_range, anchor_width, anchor_height = _drawing_anchor(getattr(chart, "anchor", None))
        preview, metadata = _render_chart_preview(chart, cached_workbook)
        visual_id = f"excel:{worksheet.title}:chart-{chart_offset + chart_index}"
        assets.append(
            VisualAsset(
                visual_id=visual_id,
                kind="chart_preview" if preview else "unrendered_chart",
                source=SourcePointer(
                    source_type="excel",
                    file_name=source_file,
                    sheet_name=worksheet.title,
                    cell=anchor_cell,
                    parser="openpyxl-chart-preview",
                    extraction_confidence=0.9 if preview else 1.0,
                ),
                media_type="image/png" if preview else None,
                width_px=anchor_width,
                height_px=anchor_height,
                sha256=hashlib.sha256(preview).hexdigest() if preview else None,
                delivery_status="ready_for_vision" if preview else "rendering_unavailable",
                metadata={"anchor_range": anchor_range, **metadata},
                image_bytes=preview,
            )
        )
    return assets


def _xlsx_package_media_assets(
    content: bytes,
    *,
    file_name: str,
    asset_offset: int,
    asset_limit: int,
    known_hashes: set[str],
) -> list[VisualAsset]:
    """Record Office media that ``openpyxl`` cannot expose as images.

    Excel workbooks can contain legacy EMF/WMF drawings and other package
    media which are not supported by Pillow/openpyxl.  Dropping their
    *existence* would make the evidence JSON incomplete, even though no local
    vision model can consume their binary form.  Keep a compact, auditable
    metadata-only asset instead.  It is deliberately unanchored: guessing a
    sheet or cell from a legacy drawing relationship would be less reliable
    than stating that its location is unavailable.
    """

    if asset_limit <= 0:
        return []
    assets: list[VisualAsset] = []
    source_file = Path(file_name).name
    try:
        archive = zipfile.ZipFile(BytesIO(content))
    except (OSError, zipfile.BadZipFile):
        return assets
    with archive:
        for member in archive.infolist():
            normalized = member.filename.lower()
            if not normalized.startswith("xl/media/") or member.is_dir():
                continue
            if len(assets) >= asset_limit:
                break
            suffix = Path(member.filename).suffix.lower().lstrip(".")
            if not suffix:
                continue
            raw = archive.read(member)
            digest = hashlib.sha256(raw).hexdigest()
            if digest in known_hashes:
                continue
            media_type = {
                "emf": "image/x-emf",
                "wmf": "image/x-wmf",
                "svg": "image/svg+xml",
                "jpg": "image/jpeg",
                "jpeg": "image/jpeg",
            }.get(suffix, f"image/{suffix}")
            visual_id = f"excel:package:media-{asset_offset + len(assets) + 1}"
            assets.append(
                VisualAsset(
                    visual_id=visual_id,
                    kind="embedded_image",
                    source=SourcePointer(
                        source_type="excel",
                        file_name=source_file,
                        parser="xlsx-zip-media",
                        extraction_confidence=1.0,
                    ),
                    media_type=media_type,
                    sha256=digest,
                    delivery_status="metadata_only",
                    metadata={
                        "archive_member": member.filename,
                        "image_format": suffix,
                        "unbound_to_sheet": True,
                        "read_error": "unsupported_office_media_format",
                    },
                )
            )
    return assets


def _xlsx_populated_coordinates(worksheet: Any, cached_sheet: Any) -> dict[int, list[int]]:
    """Return only cells with a value or formula in either workbook view.

    Excel dimensions are frequently inflated by whole-column formatting,
    historic edits, or one styled cell in column XFD.  Those formatting-only
    cells contain no business evidence, and iterating their full rectangle can
    wrongly reject a small workbook.  ``openpyxl`` has already materialised
    cells in ``_cells``; use it only to select coordinates, then read each
    selected cell through the normal worksheet API.
    """

    coordinates: set[tuple[int, int]] = set()
    for sheet in (worksheet, cached_sheet):
        for coordinate, cell in getattr(sheet, "_cells", {}).items():
            if not isinstance(coordinate, tuple) or len(coordinate) != 2:
                continue
            row_index, column_index = coordinate
            if row_index >= 1 and column_index >= 1 and getattr(cell, "value", None) is not None:
                coordinates.add((row_index, column_index))
    grouped: dict[int, list[int]] = {}
    for row_index, column_index in sorted(coordinates):
        grouped.setdefault(row_index, []).append(column_index)
    return grouped


def _parse_xlsx_to_intermediate(file_name: str, content: bytes) -> IntermediateDocument:
    """Read a modern workbook without persisting it and preserve its layout."""

    try:
        formulas = load_workbook(BytesIO(content), data_only=False, read_only=False)
        cached = load_workbook(BytesIO(content), data_only=True, read_only=False)
    except Exception as exc:  # openpyxl exposes several format-specific errors
        raise FinanceIntakeError("无法读取该 Excel 文件，请确认它是未加密的 .xlsx 工作簿。") from exc

    nonempty_cells = 0
    tables: list[IntermediateTable] = []
    blocks: list[EvidenceBlock] = []
    visual_assets: list[VisualAsset] = []
    try:
        for worksheet in formulas.worksheets:
            cached_sheet = cached[worksheet.title]
            populated_coordinates = _xlsx_populated_coordinates(worksheet, cached_sheet)
            rows: list[list[IntermediateCell]] = []
            for row_index, column_indexes in populated_coordinates.items():
                cells = [
                    _make_excel_cell(
                        row_index=row_index,
                        column_index=column_index,
                        formula_value=worksheet.cell(row=row_index, column=column_index).value,
                        cached_value=cached_sheet.cell(row=row_index, column=column_index).value,
                    )
                    for column_index in column_indexes
                ]
                if _is_nonempty_row(cells):
                    nonempty_cells += sum(
                        cell.value is not None or cell.formula is not None or cell.cached_value is not None for cell in cells
                    )
                    if nonempty_cells > MAX_EXCEL_NONEMPTY_CELLS:
                        raise FinanceIntakeError(
                            f"Excel 非空单元格超过当前安全上限 {MAX_EXCEL_NONEMPTY_CELLS} 个；"
                            "请拆分文件，或由管理员根据可用内存调整 DOCUMENT_EXCEL_MAX_NONEMPTY_CELLS。"
                        )
                    rows.append(cells)
            semantic_max_row = max(populated_coordinates, default=1)
            semantic_max_column = max(
                (column for columns in populated_coordinates.values() for column in columns),
                default=1,
            )
            table = IntermediateTable(
                table_id=f"excel:{worksheet.title}",
                title=worksheet.title,
                sheet_name=worksheet.title,
                sheet_state=worksheet.sheet_state,
                range=f"A1:{get_column_letter(semantic_max_column)}{semantic_max_row}",
                parser="openpyxl-3.1",
                extraction_confidence=1.0,
                headers=_headers_for_rows(rows),
                rows=rows,
                merged_ranges=[str(item) for item in worksheet.merged_cells.ranges],
            )
            source = SourcePointer(
                source_type="excel",
                file_name=Path(file_name).name,
                sheet_name=worksheet.title,
                parser="openpyxl-3.1",
                extraction_confidence=1.0,
            )
            section_tables = split_sheet_into_tables(table, source_type="excel", file_name=Path(file_name).name)
            tables.extend(section_tables)
            remaining_visual_capacity = max(0, MAX_EXCEL_VISUAL_ASSETS - len(visual_assets))
            sheet_visual_assets = _xlsx_visual_assets_for_sheet(
                worksheet,
                cached,
                file_name=Path(file_name).name,
                asset_offset=len(visual_assets),
                asset_limit=remaining_visual_capacity,
            )
            visual_assets.extend(sheet_visual_assets)
            blocks.append(
                evidence_block(
                    block_id=f"excel:sheet:{worksheet.title}",
                    kind="worksheet",
                    source=source,
                    table_id=table.table_id,
                    metadata={
                        "sheet_state": worksheet.sheet_state,
                        "range": table.range,
                        "merged_ranges": table.merged_ranges,
                        "row_count": semantic_max_row,
                        "column_count": semantic_max_column,
                        "source_dimension": worksheet.calculate_dimension(),
                        "formatting_dimension": {"row_count": worksheet.max_row, "column_count": worksheet.max_column},
                        "populated_cell_count": sum(len(columns) for columns in populated_coordinates.values()),
                        "logical_table_count": len(section_tables),
                        "visual_asset_count": len(sheet_visual_assets),
                    },
                )
            )
            for visual in sheet_visual_assets:
                blocks.append(
                    evidence_block(
                        block_id=f"evidence:{visual.visual_id}",
                        kind="embedded_visual",
                        source=visual.source,
                        metadata={
                            "visual_id": visual.visual_id,
                            "kind": visual.kind,
                            "media_type": visual.media_type,
                            "width_px": visual.width_px,
                            "height_px": visual.height_px,
                            "delivery_status": visual.delivery_status,
                            **visual.metadata,
                        },
                    )
                )
            for section in section_tables:
                blocks.append(
                    evidence_block(
                        block_id=f"evidence:{section.table_id}",
                        kind="table",
                        source=source.model_copy(update={"section_id": section.section_id}),
                        table_id=section.table_id,
                        metadata={
                            "parent_table_id": section.parent_table_id,
                            "range": section.range,
                            "header_range": section.header_range,
                            "headers": section.headers,
                            "scope": section.scope,
                            "scope_confidence": section.scope_confidence,
                            "merged_ranges": section.merged_ranges,
                        },
                    )
                )
        remaining_visual_capacity = max(0, MAX_EXCEL_VISUAL_ASSETS - len(visual_assets))
        package_media_assets = _xlsx_package_media_assets(
            content,
            file_name=Path(file_name).name,
            asset_offset=len(visual_assets),
            asset_limit=remaining_visual_capacity,
            known_hashes={asset.sha256 for asset in visual_assets if asset.sha256},
        )
        visual_assets.extend(package_media_assets)
        for visual in package_media_assets:
            blocks.append(
                evidence_block(
                    block_id=f"evidence:{visual.visual_id}",
                    kind="embedded_visual",
                    source=visual.source,
                    metadata={
                        "visual_id": visual.visual_id,
                        "kind": visual.kind,
                        "media_type": visual.media_type,
                        "width_px": visual.width_px,
                        "height_px": visual.height_px,
                        "delivery_status": visual.delivery_status,
                        **visual.metadata,
                    },
                )
            )
    finally:
        formulas.close()
        cached.close()
    _annotate_large_tabular_document(blocks, nonempty_cells)
    return finalize_intermediate_evidence(
        IntermediateDocument(
            source_type="excel",
            file_name=Path(file_name).name,
            parser="openpyxl-3.1",
            tables=tables,
            blocks=blocks,
            visual_assets=visual_assets,
        ),
        content,
    )


def _xls_sheet_state(visibility: Any) -> str:
    return {0: "visible", 1: "hidden", 2: "veryHidden"}.get(visibility, "visible")


def _xls_merged_range(row_low: int, row_high: int, column_low: int, column_high: int) -> str:
    return f"{get_column_letter(column_low + 1)}{row_low + 1}:{get_column_letter(column_high)}{row_high}"


def _parse_xls_to_intermediate(file_name: str, content: bytes) -> IntermediateDocument:
    """Read legacy BIFF ``.xls`` workbooks with xlrd in memory only."""

    try:
        import xlrd
    except ImportError as exc:
        raise FinanceIntakeError("旧版 .xls 解析依赖未安装，请安装 xlrd 后重试。") from exc
    try:
        workbook = xlrd.open_workbook(file_contents=content, on_demand=False, formatting_info=False)
    except Exception as exc:
        raise FinanceIntakeError("无法读取该 .xls 文件，请确认它未损坏、未加密且为 Excel 97-2003 格式。") from exc

    nonempty_cells = 0
    tables: list[IntermediateTable] = []
    blocks: list[EvidenceBlock] = []
    for worksheet in workbook.sheets():
        if worksheet.nrows * worksheet.ncols > MAX_EXCEL_GRID_CELLS:
            raise FinanceIntakeError(
                f"工作表“{worksheet.name}”的使用区域超过 {MAX_EXCEL_GRID_CELLS} 个单元格，请拆分后上传。"
            )
        rows: list[list[IntermediateCell]] = []
        for row_index in range(worksheet.nrows):
            cells: list[IntermediateCell] = []
            for column_index in range(worksheet.ncols):
                source_cell = worksheet.cell(row_index, column_index)
                value: Any = source_cell.value
                if source_cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        value = xlrd.xldate_as_datetime(source_cell.value, workbook.datemode)
                    except (ValueError, OverflowError):
                        value = source_cell.value
                display_text = _display(value)
                cells.append(
                    IntermediateCell(
                        coordinate=f"{get_column_letter(column_index + 1)}{row_index + 1}",
                        row_index=row_index + 1,
                        column_index=column_index + 1,
                        value=_serialise_cell_value(value),
                        display_text=display_text,
                        cached_value=_serialise_cell_value(value),
                        numeric_candidates=numeric_candidates_for_text(display_text),
                    )
                )
            if _is_nonempty_row(cells):
                nonempty_cells += sum(cell.value is not None for cell in cells)
                if nonempty_cells > MAX_EXCEL_NONEMPTY_CELLS:
                    raise FinanceIntakeError(
                        f"Excel 非空单元格超过当前安全上限 {MAX_EXCEL_NONEMPTY_CELLS} 个；"
                        "请拆分文件，或由管理员根据可用内存调整 DOCUMENT_EXCEL_MAX_NONEMPTY_CELLS。"
                    )
                rows.append(cells)
        table = IntermediateTable(
            table_id=f"xls:{worksheet.name}",
            title=worksheet.name,
            sheet_name=worksheet.name,
            sheet_state=_xls_sheet_state(getattr(worksheet, "visibility", 0)),
            range=f"A1:{get_column_letter(max(worksheet.ncols, 1))}{max(worksheet.nrows, 1)}",
            parser="xlrd-2",
            extraction_confidence=1.0,
            headers=_headers_for_rows(rows),
            rows=rows,
            merged_ranges=[_xls_merged_range(*merged) for merged in getattr(worksheet, "merged_cells", [])],
        )
        source = SourcePointer(
            source_type="excel",
            file_name=Path(file_name).name,
            sheet_name=worksheet.name,
            parser="xlrd-2",
            extraction_confidence=1.0,
        )
        section_tables = split_sheet_into_tables(table, source_type="excel", file_name=Path(file_name).name)
        tables.extend(section_tables)
        blocks.append(
            evidence_block(
                block_id=f"xls:sheet:{worksheet.name}",
                kind="worksheet",
                source=source,
                table_id=table.table_id,
                metadata={
                    "sheet_state": table.sheet_state,
                    "range": table.range,
                    "merged_ranges": table.merged_ranges,
                    "row_count": worksheet.nrows,
                    "column_count": worksheet.ncols,
                    "logical_table_count": len(section_tables),
                },
            )
        )
        for section in section_tables:
            blocks.append(
                evidence_block(
                    block_id=f"evidence:{section.table_id}",
                    kind="table",
                    source=source.model_copy(update={"section_id": section.section_id}),
                    table_id=section.table_id,
                    metadata={
                        "parent_table_id": section.parent_table_id,
                        "range": section.range,
                        "header_range": section.header_range,
                        "headers": section.headers,
                        "scope": section.scope,
                        "scope_confidence": section.scope_confidence,
                        "merged_ranges": section.merged_ranges,
                    },
                )
            )
    _annotate_large_tabular_document(blocks, nonempty_cells)
    return finalize_intermediate_evidence(
        IntermediateDocument(
            source_type="excel",
            file_name=Path(file_name).name,
            parser="xlrd-2",
            tables=tables,
            blocks=blocks,
        ),
        content,
    )


def _decode_csv(content: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise FinanceIntakeError("无法识别 CSV 编码，请另存为 UTF-8 或 GBK/GB18030 后重试。")


def _parse_csv_to_intermediate(file_name: str, content: bytes) -> IntermediateDocument:
    """Parse a UTF-8/GBK CSV as one auditable, in-memory table."""

    text = _decode_csv(content)
    try:
        dialect = csv.Sniffer().sniff(text[:8_192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(StringIO(text), dialect)
    rows: list[list[IntermediateCell]] = []
    nonempty_cells = 0
    max_columns = 0
    grid_cells = 0
    for row_index, values in enumerate(reader, start=1):
        grid_cells += len(values)
        if grid_cells > MAX_EXCEL_GRID_CELLS:
            raise FinanceIntakeError(f"CSV 使用区域超过 {MAX_EXCEL_GRID_CELLS} 个单元格，请拆分后上传。")
        cells = [
            IntermediateCell(
                coordinate=f"{get_column_letter(column_index)}{row_index}",
                row_index=row_index,
                column_index=column_index,
                value=value if value != "" else None,
                display_text=value or None,
                cached_value=value if value != "" else None,
                numeric_candidates=numeric_candidates_for_text(value),
            )
            for column_index, value in enumerate(values, start=1)
        ]
        if _is_nonempty_row(cells):
            nonempty_cells += sum(cell.value is not None for cell in cells)
            if nonempty_cells > MAX_EXCEL_NONEMPTY_CELLS:
                raise FinanceIntakeError(f"CSV 非空单元格超过 {MAX_EXCEL_NONEMPTY_CELLS} 个，请拆分后上传。")
            rows.append(cells)
        max_columns = max(max_columns, len(cells))
    if not rows:
        raise FinanceIntakeError("CSV 中没有可读取的数据行。")
    table_name = Path(file_name).stem or "CSV"
    table = IntermediateTable(
        table_id=f"csv:{table_name}",
        title=table_name,
        sheet_name=table_name,
        sheet_state="visible",
        range=f"A1:{get_column_letter(max(max_columns, 1))}{len(rows)}",
        parser="csv",
        headers=_headers_for_rows(rows),
        rows=rows,
    )
    source = SourcePointer(
        source_type="excel",
        file_name=Path(file_name).name,
        sheet_name=table_name,
        parser="csv",
        extraction_confidence=1.0,
    )
    section_tables = split_sheet_into_tables(table, source_type="excel", file_name=Path(file_name).name)
    blocks = [
        evidence_block(
            block_id=f"csv:sheet:{table_name}",
            kind="worksheet",
            source=source,
            table_id=table.table_id,
            metadata={"sheet_state": "visible", "range": table.range, "row_count": len(rows), "column_count": max_columns},
        ),
        *[
            evidence_block(
                block_id=f"evidence:{section.table_id}",
                kind="table",
                source=source.model_copy(update={"section_id": section.section_id}),
                table_id=section.table_id,
                metadata={
                    "parent_table_id": section.parent_table_id,
                    "range": section.range,
                    "header_range": section.header_range,
                    "headers": section.headers,
                    "scope": section.scope,
                    "scope_confidence": section.scope_confidence,
                },
            )
            for section in section_tables
        ],
    ]
    _annotate_large_tabular_document(blocks, nonempty_cells)
    return finalize_intermediate_evidence(
        IntermediateDocument(
            source_type="excel",
            file_name=Path(file_name).name,
            parser="csv",
            tables=section_tables,
            blocks=blocks,
        ),
        content,
    )


def parse_excel_to_intermediate(file_name: str, content: bytes) -> IntermediateDocument:
    """Read XLSX, legacy XLS, or CSV without persisting customer files."""

    suffix = _validate_tabular_upload(file_name, content)
    if suffix == ".xlsx":
        return _parse_xlsx_to_intermediate(file_name, content)
    if suffix == ".xls":
        return _parse_xls_to_intermediate(file_name, content)
    return _parse_csv_to_intermediate(file_name, content)


def _header_index(row: list[IntermediateCell]) -> dict[str, int]:
    result: dict[str, int] = {}
    for index, cell in enumerate(row):
        label = _normalised_label(cell.display_text or "")
        if label:
            result[label] = index
    return result


def _first_table_row(table: IntermediateTable, predicate: Any) -> tuple[int, list[IntermediateCell]] | None:
    for index, row in enumerate(table.rows):
        if predicate(row):
            return index, row
    return None


def _section_title_before_header(table: IntermediateTable, header_row_index: int) -> str | None:
    """Return the nearest preceding title-like row for an independent block."""

    for row in reversed(table.rows[:header_row_index]):
        values = [(cell.display_text or "").strip() for cell in row]
        nonempty = [value for value in values if value]
        if not nonempty:
            continue
        # A title row is usually a single populated (often merged) cell.  Do
        # not mistake the preceding data row for a section label.
        if len(nonempty) == 1:
            return nonempty[0]
        return None
    return None


def _standardize_wide_table(
    document: IntermediateDocument,
    table: IntermediateTable,
    semantic_index: dict[str, SemanticLabelMapping],
    default_year: str | None,
) -> list[StandardFinancialFact]:
    header_match = _first_table_row(
        table,
        lambda row: any(_header_role(cell.display_text, semantic_index) == "__PERIOD__" for cell in row),
    )
    if header_match is None:
        return []
    header_row_index, header = header_match
    period_index = next(
        index
        for index, cell in enumerate(header)
        if _header_role(cell.display_text, semantic_index) == "__PERIOD__"
    )
    metrics_by_column = {
        index: resolution
        for index, cell in enumerate(header)
        if (resolution := _metric_resolution(cell.display_text or "", semantic_index)) is not None
    }
    if not metrics_by_column:
        return []

    facts: list[StandardFinancialFact] = []
    section_index = 1
    section_title = table.title or _section_title_before_header(table, header_row_index)
    section_id = table.section_id or f"{table.table_id}:section-{section_index}"
    for row_index, row in enumerate(table.rows[header_row_index + 1 :], start=header_row_index + 1):
        # Spreadsheet exports commonly place multiple independently headed
        # data blocks on one sheet.  Preserve their boundary so the UI can ask
        # the customer which block is authoritative instead of merging them.
        if period_index < len(row) and _header_role(row[period_index].display_text, semantic_index) == "__PERIOD__":
            section_index += 1
            section_title = _section_title_before_header(table, row_index)
            section_id = (
                f"{table.section_id}:subsection-{section_index}"
                if table.section_id
                else f"{table.table_id}:section-{section_index}"
            )
            continue
        period_cell = row[period_index] if period_index < len(row) else None
        period_value = period_cell.cached_value if period_cell and period_cell.cached_value is not None else period_cell.value if period_cell else None
        period = _period_from_value(period_value, default_year=default_year)
        if period is None:
            continue
        for column_index, (metric, mapping_method, confidence) in metrics_by_column.items():
            if column_index >= len(row):
                continue
            cell = row[column_index]
            value = cell.cached_value if cell.cached_value is not None else cell.value
            # A blank cell is not a zero.  Do not turn it into an invalid
            # financial fact merely because its column has a valid mapping.
            if _number(value) is None:
                continue
            facts.append(
                _build_fact(
                    metric,
                    header[column_index].display_text or metric.name,
                    value,
                    period,
                    SourcePointer(
                        source_type=document.source_type,
                        file_name=document.file_name,
                        sheet_name=table.sheet_name,
                        page_number=table.page_number,
                        cell=cell.coordinate,
                        row_index=cell.row_index,
                        column_index=cell.column_index,
                        section_id=section_id,
                        section_title=section_title,
                        parser=table.parser,
                        extraction_confidence=table.extraction_confidence,
                    ),
                    mapping_method=mapping_method,
                    confidence=confidence,
                )
            )
    return facts


def _statement_header_row(table: IntermediateTable, semantic_index: dict[str, SemanticLabelMapping]) -> tuple[int, list[IntermediateCell]] | None:
    """Find a statement header, including a PDF header collapsed into one cell.

    Some valid PDFs draw no vertical grid lines in their heading row.  Table
    extractors then return ``项目\n附注\n2024年\n2023年`` in the first cell even
    though subsequent rows have four columns.  Recover only this explicit,
    line-delimited structure and retain the original coordinates of the blank
    heading cells; the canonical table evidence is never rewritten.
    """

    direct = _first_table_row(
        table,
        lambda row: any(_header_role(cell.display_text, semantic_index) == "__ITEM__" for cell in row),
    )
    if direct is not None:
        return direct
    for index, row in enumerate(table.rows):
        nonempty = [cell for cell in row if (cell.display_text or "").strip()]
        if len(nonempty) != 1 or len(row) < 3:
            continue
        labels = [label.strip() for label in re.split(r"[\r\n]+", nonempty[0].display_text or "") if label.strip()]
        if len(labels) < 3 or len(labels) > len(row):
            continue
        if not any(_header_role(label, semantic_index) == "__ITEM__" for label in labels):
            continue
        recovered = [
            cell.model_copy(update={"display_text": labels[column_index] if column_index < len(labels) else cell.display_text})
            for column_index, cell in enumerate(row)
        ]
        return index, recovered
    return None


def _statement_amount_columns(
    header: list[IntermediateCell],
    semantic_index: dict[str, SemanticLabelMapping],
) -> tuple[int | None, int | None, str | None, str | None]:
    """Locate amount columns without assuming one PDF's header wording."""

    current_index = next(
        (index for index, cell in enumerate(header) if _header_role(cell.display_text, semantic_index) == "__CURRENT_AMOUNT__"),
        None,
    )
    prior_index = next(
        (index for index, cell in enumerate(header) if _header_role(cell.display_text, semantic_index) == "__PRIOR_AMOUNT__"),
        None,
    )
    dated_columns = [
        (index, period)
        for index, cell in enumerate(header)
        if (period := _period_from_value(cell.display_text)) is not None
    ]
    if current_index is None and dated_columns:
        current_index = dated_columns[0][0]
    if prior_index is None:
        prior_index = next((index for index, _period in dated_columns if index != current_index), None)
    current_period = _period_from_value(header[current_index].display_text) if current_index is not None else None
    prior_period = _period_from_value(header[prior_index].display_text) if prior_index is not None else _previous_period(current_period)
    return current_index, prior_index, current_period, prior_period


def _standardize_statement_table(
    document: IntermediateDocument,
    table: IntermediateTable,
    semantic_index: dict[str, SemanticLabelMapping],
) -> list[StandardFinancialFact]:
    header_match = _statement_header_row(table, semantic_index)
    if header_match is None:
        return []
    header_row_index, header = header_match
    item_index = next((index for index, cell in enumerate(header) if _header_role(cell.display_text, semantic_index) == "__ITEM__"), None)
    if item_index is None:
        return []
    # PDF tables often merge the period into a header such as “2025年本期
    # 金额”; retain that period while still recognising the amount column.
    current_index, prior_index, current_period, prior_period = _statement_amount_columns(header, semantic_index)
    if current_index is None and prior_index is None:
        return []
    facts: list[StandardFinancialFact] = []
    for row in table.rows[header_row_index + 1 :]:
        if item_index >= len(row):
            continue
        item_cell = row[item_index]
        resolution = _metric_resolution(item_cell.display_text or "", semantic_index)
        if resolution is None:
            continue
        metric, mapping_method, confidence = resolution
        for value_index, period in ((current_index, current_period), (prior_index, prior_period)):
            if value_index is None or value_index >= len(row):
                continue
            value_cell = row[value_index]
            value = value_cell.cached_value if value_cell.cached_value is not None else value_cell.value
            # Table extractors sometimes leave a subordinate label (for
            # example "of which: cost of sales") in a monetary column, or a
            # non-calculated formula surface.  The original cell remains in
            # Evidence JSON, but it is not a numeric financial fact.  Do not
            # manufacture a ``value=None`` fact merely because the row label
            # maps to a known metric.
            if _number(value) is None:
                continue
            facts.append(
                _build_fact(
                    metric,
                    item_cell.display_text or metric.name,
                    value,
                    period,
                    SourcePointer(
                        source_type=document.source_type,
                        file_name=document.file_name,
                        sheet_name=table.sheet_name,
                        page_number=table.page_number,
                        cell=value_cell.coordinate,
                        row_index=value_cell.row_index,
                        column_index=value_cell.column_index,
                        section_id=table.section_id,
                        section_title=table.title,
                        parser=table.parser,
                        extraction_confidence=table.extraction_confidence,
                    ),
                    mapping_method=mapping_method,
                    confidence=confidence,
                )
            )
    return facts


def standardize_tables(
    document: IntermediateDocument,
    semantic_mappings: Iterable[SemanticLabelMapping] = (),
    semantic_table_contexts: Iterable[SemanticTableContext] = (),
) -> IntakeResult:
    if document.source_type not in {"excel", "pdf", "word"}:
        raise FinanceIntakeError("表格标准化器只能处理 Excel、PDF 或 Word 中间 JSON。")
    facts: list[StandardFinancialFact] = []
    semantic_mappings = list(semantic_mappings)
    semantic_context_by_table = {context.table_id: context for context in semantic_table_contexts}
    for table in document.tables:
        sheet_label = (table.sheet_name or table.title or "").lower()
        # The published customer template contains explanation, example and
        # field-dictionary sheets.  Preserve them in intermediate JSON but do
        # not silently ingest them as the customer's actual financial facts.
        if document.source_type == "excel" and any(token.lower() in sheet_label for token in NON_INGESTION_SHEET_TOKENS):
            continue
        semantic_index = _semantic_mapping_index(semantic_mappings, table.table_id)
        table_context = semantic_context_by_table.get(table.table_id)
        facts.extend(_standardize_wide_table(document, table, semantic_index, table_context.period_year if table_context else None))
        facts.extend(_standardize_statement_table(document, table, semantic_index))
    issues: list[ValidationIssue] = []
    large_workbook_blocks = [
        block
        for block in document.blocks
        if block.kind == "worksheet" and block.metadata.get("large_workbook_mode") is True
    ]
    if large_workbook_blocks:
        populated_cell_count = max(
            int(block.metadata.get("workbook_populated_cell_count", 0)) for block in large_workbook_blocks
        )
        _issue(
            issues,
            "warning",
            "large_excel_chunked",
            (
                f"工作簿包含 {populated_cell_count} 个非空单元格，已完整保留为可追溯 Evidence；"
                f"模型输入按最多 {MAX_MODEL_CONTEXT_CHUNK_CHARACTERS} 字符的来源关联片段提供。"
            ),
            large_workbook_blocks[0].source,
        )
    if not facts:
        _issue(issues, "warning", "no_mapped_financial_facts", "未识别到可确定映射的财务字段；已保留中间 JSON，需补充字段映射或客户确认。")
    _validate_facts(facts, issues)
    return IntakeResult(intermediate=document, standard=StandardFinancialDocument(facts=facts), validation=issues)


def standardize_excel(document: IntermediateDocument) -> IntakeResult:
    if document.source_type != "excel":
        raise FinanceIntakeError("Excel 标准化器只能处理 Excel 中间 JSON。")
    return standardize_tables(document)


_CLAUSE_SPLIT = re.compile(r"[，,。；;\n]+")
_AMOUNT = re.compile(r"(?P<amount>[-+]?\(?\d+(?:,\d{3})*(?:\.\d+)?\)?)\s*(?P<unit>亿元|万元|元|人民币)?")


def _amount_in_yuan(value: str, unit: str | None, default_unit: Literal["yuan", "ten_thousand_yuan"] | None) -> float | None:
    number = _number(value)
    if number is None:
        return None
    factors = {"元": 1, "人民币": 1, "万元": 10_000, "亿元": 100_000_000, "yuan": 1, "ten_thousand_yuan": 10_000}
    chosen = unit or default_unit
    return number * factors[chosen] if chosen in factors else None


def _period_and_company(text: str, explicit_period: str | None) -> tuple[str | None, str | None]:
    period = explicit_period or _period_from_value(text)
    company_match = re.search(r"(?P<name>[\u4e00-\u9fffA-Za-z0-9（）()]{2,30}(?:公司|集团))", text)
    return period, company_match.group("name") if company_match else None


def parse_text_to_intermediate(
    text: str,
    *,
    file_name: str | None = None,
    source_bytes: bytes | None = None,
) -> IntermediateDocument:
    """Turn typed text or a decoded ``.txt`` file into canonical evidence.

    ``source_bytes`` lets a text-file upload retain the hash and byte size of
    the actual uploaded file. Typed chat input deliberately falls back to the
    normalized UTF-8 text, because it has no separate binary upload.
    """

    clean_text = text.strip()
    if not clean_text:
        raise FinanceIntakeError("文字描述不能为空。")
    if len(clean_text) > MAX_TEXT_CHARACTERS:
        raise FinanceIntakeError(f"文字描述超过 {MAX_TEXT_CHARACTERS} 字限制。")
    cell = IntermediateCell(
        coordinate="text:1",
        row_index=1,
        column_index=1,
        value=clean_text,
        display_text=clean_text,
        numeric_candidates=numeric_candidates_for_text(clean_text),
    )
    source = SourcePointer(
        source_type="text",
        file_name=Path(file_name).name if file_name else None,
        section_id="text:input",
        text_span=(0, len(clean_text)),
        parser="rule-based-text-v1",
        extraction_confidence=1.0,
    )
    blocks: list[EvidenceBlock] = []
    cursor = 0
    for sequence, raw_line in enumerate(clean_text.splitlines(keepends=True), start=1):
        line = raw_line.strip()
        line_start = cursor + (len(raw_line) - len(raw_line.lstrip()))
        cursor += len(raw_line)
        if not line:
            continue
        line_end = line_start + len(line)
        blocks.append(
            evidence_block(
                block_id=f"text:paragraph:{len(blocks) + 1}",
                kind="paragraph",
                source=SourcePointer(
                    source_type="text",
                    file_name=Path(file_name).name if file_name else None,
                    section_id=f"text:paragraph:{len(blocks) + 1}",
                    text_span=(line_start, line_end),
                    parser="rule-based-text-v1",
                    extraction_confidence=1.0,
                ),
                text=line,
                metadata={"sequence": sequence},
            )
        )
    if not blocks:
        blocks.append(
            evidence_block(
                block_id="text:paragraph:1",
                kind="paragraph",
                source=source,
                text=clean_text,
                metadata={"sequence": 1},
            )
        )
    return finalize_intermediate_evidence(
        IntermediateDocument(
            source_type="text",
            file_name=Path(file_name).name if file_name else None,
            parser="rule-based-text-v1",
            raw_text=clean_text,
            tables=[IntermediateTable(table_id="text:input", title="文字描述", headers=["原始文字"], rows=[[cell]])],
            blocks=blocks,
        ),
        source_bytes if source_bytes is not None else clean_text.encode("utf-8"),
    )


def standardize_text(
    document: IntermediateDocument,
    *,
    period: str | None = None,
    default_unit: Literal["yuan", "ten_thousand_yuan"] | None = None,
) -> IntakeResult:
    if document.source_type not in {"text", "word"} or not document.raw_text:
        raise FinanceIntakeError("文字标准化器只能处理文字或 Word 中间 JSON。")
    initial_period, company_name = _period_and_company(document.raw_text, period)
    active_period = initial_period
    last_metric: MetricDefinition | None = None
    facts: list[StandardFinancialFact] = []
    issues: list[ValidationIssue] = []
    offset = 0
    for clause in _CLAUSE_SPLIT.split(document.raw_text):
        start = document.raw_text.find(clause, offset)
        offset = start + len(clause) if start >= 0 else offset
        if not clause.strip():
            continue
        explicit_clause_period = _period_from_value(clause)
        if explicit_clause_period:
            active_period = explicit_clause_period
        is_previous = any(term in clause for term in ("去年", "上年", "同期", "上期"))
        alias_match = next(
            ((item, alias, clause.find(alias)) for item in METRICS for alias in item.aliases if alias in clause),
            None,
        )
        metric = alias_match[0] if alias_match is not None else None
        if metric is not None:
            last_metric = metric
        elif is_previous:
            metric = last_metric
        if metric is None:
            continue
        amount_segment = clause[alias_match[2] + len(alias_match[1]) :] if alias_match is not None else clause
        amount_match = _AMOUNT.search(amount_segment)
        if amount_match is None:
            _issue(
                issues,
                "warning",
                "amount_missing",
                f"已识别“{metric.name}”，但未找到金额。",
                SourcePointer(
                    source_type=document.source_type,
                    file_name=document.file_name,
                    text_span=(max(start, 0), max(start, 0) + len(clause)),
                ),
            )
            continue
        target_period = _previous_period(active_period) if is_previous else active_period
        converted = _amount_in_yuan(amount_match.group("amount"), amount_match.group("unit"), default_unit)
        source = SourcePointer(
            source_type=document.source_type,
            file_name=document.file_name,
            text_span=(max(start, 0), max(start, 0) + len(clause)),
        )
        fact = _build_fact(
            metric,
            metric.name,
            converted,
            target_period,
            source,
            confidence=0.95,
            unit="yuan" if converted is not None else "unknown",
        )
        if converted is None:
            _issue(issues, "warning", "unit_missing", f"{metric.name} 的金额缺少元、万元或亿元单位，需客户确认。", source)
        facts.append(fact)
    if not facts:
        _issue(issues, "warning", "no_mapped_financial_facts", "未从文字中识别到可确定映射的财务指标；请明确期间、科目、金额和单位。")
    _validate_facts(facts, issues)
    return IntakeResult(
        intermediate=document,
        standard=StandardFinancialDocument(company_name=company_name, facts=facts),
        validation=issues,
    )


def ingest_text(
    text: str,
    *,
    period: str | None = None,
    default_unit: Literal["yuan", "ten_thousand_yuan"] | None = None,
    file_name: str | None = None,
    source_bytes: bytes | None = None,
) -> IntakeResult:
    return standardize_text(
        parse_text_to_intermediate(text, file_name=file_name, source_bytes=source_bytes),
        period=period,
        default_unit=default_unit,
    )


def ingest_text_file(file_name: str, content: bytes) -> IntakeResult:
    """Ingest a safe plain-text upload through the same path as typed text.

    UTF-8 (including BOM) is preferred.  GB18030 is a reversible fallback for
    common Chinese Windows text files; other binary/unknown encodings are not
    guessed, so a customer is never shown silently corrupted evidence.
    """

    if Path(file_name).suffix.lower() != ".txt":
        raise FinanceIntakeError("文本文件入口仅接收 .txt 文件。")
    if not content:
        raise FinanceIntakeError("上传的文本文件为空。")
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = content.decode("gb18030")
        except UnicodeDecodeError as exc:
            raise FinanceIntakeError("无法识别文本文件编码；请保存为 UTF-8 或 GB18030 后重试。") from exc
    return ingest_text(text, file_name=file_name, source_bytes=content)


def ingest_excel(file_name: str, content: bytes) -> IntakeResult:
    return standardize_excel(parse_excel_to_intermediate(file_name, content))


def intake_contract() -> dict[str, Any]:
    return {
        "version": "finance-intake-v1.7",
        "evidence_versions": {
            "default": "v1",
            "available": ["v1", "v2"],
            "v2_usage": "统一文件接口增加 evidence_version=v2 后返回类型化 location、全局 element_id、独立文件/快照哈希；V1 保持不变。",
        },
        "supported_now": [
            "text",
            "xlsx",
            "xls",
            "csv",
            "text_pdf",
            "docx",
            "image_vision_candidate",
            "scanned_pdf_vision_candidate",
            "zip_tabular_archive",
            "html_inline_xbrl_structure",
            "xml_repeating_records",
            "sec_submission_manifest",
        ],
        "planned_not_enabled": ["deterministic_scanned_pdf_ocr", "deterministic_image_table_ocr"],
        "storage_policy": "文字和 Excel 仅在内存解析；PDF 使用请求级临时目录调用版面/表格工具，响应结束后删除，不做持久化保存。",
        "semantic_mapping": {
            "mode": "local_qwen3_candidate_when_rule_mapping_is_empty",
            "policy": "模型只映射上传表格中已有标签；金额、期间、单元格坐标仍由程序读取，所有语义映射均须客户确认。",
        },
        "two_layer_output": {
            "intermediate": "Evidence JSON：保留原始文字、页/段/表块、Excel 单元格、工作表状态、公式、合并区域、数值表面候选与来源定位。",
            "model_context": "Evidence JSON 的有界结构化渲染，仅供模型阅读；财务取数和审计始终回到原始坐标。",
            "standard": "仅包含规则可确定映射的标准财务指标；低置信度或缺失信息进入校验结果。",
            "visual_evidence": "XLSX 内嵌图片和常见原生图表会保留为与 Evidence JSON 并列的视觉证据。图表预览仅供本地视觉模型识别趋势、标题和图例，不能生成或覆盖财务金额；视觉结论必须复核。",
        },
        "visual_evidence": {
            "supported_now": [
                "xlsx_embedded_images",
                "xlsx_bar_line_area_scatter_pie_doughnut_chart_previews",
                "image_generic_vision_candidates",
                "scanned_pdf_page_generic_vision_candidates",
            ],
            "limits": {"max_assets_per_upload": MAX_EXCEL_VISUAL_ASSETS, "max_assets_sent_to_vision": 4},
            "deferred": ["xls_legacy_drawings", "Office_shapes_and_SmartArt", "unsupported_native_chart_renderer"],
        },
        "required_for_reliable_fact": ["财务科目", "金额", "单位", "报告期间"],
        "standard_metrics": [
            {"code": metric.code, "name": metric.name, "statement_type": metric.statement_type, "aliases": list(metric.aliases)}
            for metric in METRICS
        ],
    }
