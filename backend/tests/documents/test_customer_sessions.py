from __future__ import annotations

import unittest
from io import BytesIO
from unittest.mock import patch

from backend.documents.customer_sessions import (
    _estimate_text_tokens,
    _semantic_rerank_customer_windows,
    add_files,
    bind_session_owner,
    delete_session,
    get_session,
    get_visual_asset,
    issue_customer_visual_ticket,
    retrieve,
    temporary_visual_files,
)


class CustomerDocumentSessionTests(unittest.TestCase):
    def test_owner_bound_session_cannot_be_read_appended_or_deleted_by_another_owner(self) -> None:
        owner_a = "browser:" + "a" * 64
        owner_b = "browser:" + "b" * 64
        result = add_files([("private.txt", b"private evidence")], owner_id=owner_a)
        session_id = result["session_id"]
        self.sessions = [session_id]

        self.assertIsNotNone(get_session(session_id, owner_id=owner_a))
        self.assertIsNone(get_session(session_id, owner_id=owner_b))
        with self.assertRaisesRegex(PermissionError, "owner_mismatch"):
            add_files([("foreign.txt", b"foreign")], session_id=session_id, owner_id=owner_b)
        self.assertFalse(delete_session(session_id, owner_id=owner_b))
        self.assertIsNotNone(get_session(session_id, owner_id=owner_a))

        with bind_session_owner(owner_b):
            self.assertIsNone(get_session(session_id))
            self.assertEqual(retrieve(session_id, "private evidence")["status"], "session_not_found")
        with bind_session_owner(owner_a):
            self.assertIsNotNone(get_session(session_id))
        self.assertTrue(delete_session(session_id, owner_id=owner_a))
        self.sessions = []

    def test_visual_ticket_is_scoped_and_allows_headerless_image_fetch(self) -> None:
        from PIL import Image

        owner = "browser:" + "c" * 64
        stream = BytesIO()
        Image.new("RGB", (80, 60), color=(210, 220, 230)).save(stream, format="PNG")
        result = add_files([("private.png", stream.getvalue())], owner_id=owner)
        session_id = result["session_id"]
        self.sessions = [session_id]
        with bind_session_owner(owner):
            session = get_session(session_id)
            self.assertIsNotNone(session)
            assert session is not None
            document = session.documents[0]
            visual = document.visuals[0]
            ticket = issue_customer_visual_ticket(session_id, document.document_id, visual.visual_id)
        self.assertTrue(ticket)
        self.assertIsNone(
            get_visual_asset(session_id, document.document_id, visual.visual_id, owner_id="browser:" + "d" * 64)
        )
        ticketed = get_visual_asset(
            session_id,
            document.document_id,
            visual.visual_id,
            access_ticket=ticket,
        )
        self.assertIsNotNone(ticketed)
        self.assertIsNone(
            get_visual_asset(session_id, "another-document", visual.visual_id, access_ticket=ticket)
        )
        self.assertTrue(delete_session(session_id, owner_id=owner))
        self.sessions = []

    def test_generic_global_overview_uses_structure_sampling_without_loading_reranker(self) -> None:
        candidates = [
            {
                "text": f"candidate {index}",
                "score": float(20 - index),
                "order": index,
                "is_document_index": False,
            }
            for index in range(12)
        ]
        with patch.dict("os.environ", {"CUSTOMER_ATTACHMENT_SEMANTIC_RERANK": "1"}), patch(
            "backend.sales.dense_retrieval.retrieval_runtime_status",
            return_value={"effective_device": "cpu", "generation_gpu_reserved": True},
        ), patch("backend.sales.dense_retrieval.get_reranker_model") as get_model:
            ranked, audit = _semantic_rerank_customer_windows(
                "总结附件",
                candidates,
                global_document_question=True,
            )

        self.assertIs(ranked, candidates)
        self.assertEqual(audit["status"], "structure_aware_global_sampling")
        self.assertEqual(audit["reason"], "generic_overview_has_no_semantic_target")
        get_model.assert_not_called()

    def tearDown(self) -> None:
        for session_id in getattr(self, "sessions", []):
            delete_session(session_id)

    def test_two_text_files_are_retrievable_with_provenance(self) -> None:
        result = add_files(
            [
                ("施工说明.txt", "窗洞口采用锚固安装并设置防水收口。".encode("utf-8")),
                ("项目备注.txt", "项目位于济南，外墙采用保温装饰一体板。".encode("utf-8")),
            ]
        )
        self.sessions = [result["session_id"]]
        selected = retrieve(result["session_id"], "窗洞口怎么安装？")
        self.assertEqual(selected["status"], "ok")
        self.assertEqual(len(selected["documents"]), 2)
        self.assertTrue(any(item["document_name"] == "施工说明.txt" for item in selected["evidence"]))
        self.assertTrue(selected["input_snapshot"]["canonical_evidence_preserved"])
        self.assertFalse(selected["input_snapshot"]["citation_coverage_complete"])

    def test_irrelevant_files_do_not_consume_positive_evidence_slots(self) -> None:
        relevant_lines = "\n".join(f"窗洞口锚固安装节点说明 {index}" for index in range(12))
        result = add_files(
            [
                ("施工说明.txt", relevant_lines.encode("utf-8")),
                ("项目通讯录.txt", "联系人与电话号码。".encode("utf-8")),
                ("会议通知.txt", "会议时间和会议室。".encode("utf-8")),
                ("仓库清单.txt", "办公用品库存。".encode("utf-8")),
            ]
        )
        self.sessions = [result["session_id"]]
        selected = retrieve(result["session_id"], "窗洞口如何锚固安装？")
        self.assertTrue(selected["evidence"])
        self.assertEqual(
            {item["document_name"] for item in selected["evidence"]},
            {"施工说明.txt"},
        )

    def test_valid_attachment_with_zero_overlap_expands_children_not_expired_state(self) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "经营数据"
        sheet.append(["项目", "金额"])
        sheet.append(["办公楼", 80])
        stream = BytesIO()
        workbook.save(stream)
        result = add_files([("经营数据.xlsx", stream.getvalue())])
        self.sessions = [result["session_id"]]

        selected = retrieve(result["session_id"], "请识别图片中的外墙节点")

        self.assertTrue(selected["evidence"])
        self.assertEqual(selected["evidence"][0]["evidence_scope"], "content")
        self.assertTrue(selected['navigation_evidence'])
        self.assertEqual(
            selected["input_snapshot"]["retrieval_fallback_reason"],
            "navigation_child_expansion_relevance_unverified",
        )

    def test_more_than_four_files_are_rejected(self) -> None:
        with self.assertRaisesRegex(ValueError, "maximum_four_files"):
            add_files([(f"{index}.txt", b"test") for index in range(5)])

    def test_existing_session_bytes_count_toward_total_limit(self) -> None:
        with patch("backend.documents.customer_sessions.MAX_FILE_BYTES", 1_000), patch(
            "backend.documents.customer_sessions.MAX_SESSION_BYTES", 1_000
        ):
            result = add_files([("first.txt", b"A" * 700)])
            self.sessions = [result["session_id"]]
            with self.assertRaisesRegex(ValueError, "session_files_too_large"):
                add_files(
                    [("second.txt", b"B" * 400)],
                    session_id=result["session_id"],
                )

    def test_excel_evidence_returns_sheet_and_row_range(self) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "施工方案"
        sheet.append(["项目", "安装方式"])
        sheet.append(["办公楼", "锚固安装"])
        stream = BytesIO()
        workbook.save(stream)
        result = add_files([("方案.xlsx", stream.getvalue())])
        self.sessions = [result["session_id"]]
        selected = retrieve(result["session_id"], "办公楼采用什么安装方式？")
        precise = [
            citation
            for item in selected["evidence"]
            for citation in item["citations"]
            if citation.get("sheet_name") == "施工方案"
        ]
        self.assertTrue(precise)
        self.assertTrue(any(citation.get("source_range") for citation in precise))

    def test_excel_structure_overview_lists_all_sheets_for_summary_questions(self) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        summary = workbook.active
        summary.title = "经营汇总"
        summary.append(["指标", "本期"])
        summary.append(["收入", 120])
        detail = workbook.create_sheet("项目明细")
        detail.append(["项目", "金额"])
        detail.append(["办公楼", 80])
        stream = BytesIO()
        workbook.save(stream)
        result = add_files([("经营数据.xlsx", stream.getvalue())])
        self.sessions = [result["session_id"]]

        unplanned = retrieve(result["session_id"], "先给我介绍再分析一下")
        self.assertFalse(unplanned["input_snapshot"]["global_document_question"])
        self.assertEqual(unplanned["input_snapshot"]["planner_document_scope"], "unknown")

        selected = retrieve(
            result["session_id"],
            "先给我介绍再分析一下",
            document_scope="whole_document",
        )
        rendered = "\n".join(str(item.get("text") or "") for item in [*selected["evidence"], *selected.get('navigation_evidence', [])])

        self.assertIn("STRUCTURE_OVERVIEW", rendered)
        self.assertIn("经营汇总", rendered)
        self.assertIn("项目明细", rendered)
        snapshot = selected["input_snapshot"]
        self.assertTrue(snapshot["global_document_question"])
        self.assertTrue(snapshot["all_canonical_chunks_scanned"])
        self.assertEqual(snapshot["scanned_canonical_chunk_count"], result["documents"][0]["chunk_count"])
        self.assertEqual(snapshot["selected_document_index_window_count"], 0)
        self.assertGreater(snapshot['navigation_candidate_count'], 0)
        self.assertGreater(snapshot["selected_content_window_count"], 0)
        self.assertEqual(snapshot["semantic_rerank"]["status"], "structure_aware_global_sampling")

    def test_cross_document_analysis_keeps_structure_and_content_from_each_file(self) -> None:
        from openpyxl import Workbook

        files = []
        for file_name, sheet_name, rows in (
            (
                "经营表.xlsx",
                "利润表",
                [["项目", "本期"], ["营业收入", 120], ["营业支出", 90], ["利润", 30]],
            ),
            (
                "预算表.xlsx",
                "工程预算",
                [["项目", "预算"], ["材料费", 60], ["人工费", 25], ["其他", 15]],
            ),
        ):
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = sheet_name
            for row in rows:
                sheet.append(row)
            stream = BytesIO()
            workbook.save(stream)
            files.append((file_name, stream.getvalue()))

        result = add_files(files)
        self.sessions = [result["session_id"]]
        selected = retrieve(
            result["session_id"],
            "详细分析这两个文件并分别提出优化建议",
            document_scope="cross_document",
        )

        snapshot = selected["input_snapshot"]
        self.assertTrue(snapshot["global_document_question"])
        self.assertEqual(snapshot["planner_document_scope"], "cross_document")
        self.assertEqual(snapshot["selected_document_index_window_count"], 0)
        self.assertGreaterEqual(snapshot['navigation_candidate_count'], 2)
        self.assertGreaterEqual(snapshot["selected_content_window_count"], 2)
        content_documents = {
            item["document_name"]
            for item in selected["evidence"]
            if item["evidence_scope"] == "content"
        }
        self.assertEqual(content_documents, {"经营表.xlsx", "预算表.xlsx"})

    def test_image_asset_is_retained_and_selected_for_local_vlm(self) -> None:
        from PIL import Image

        stream = BytesIO()
        Image.new("RGB", (320, 180), color=(210, 215, 220)).save(stream, format="PNG")
        result = add_files([("节点图.png", stream.getvalue())])
        self.sessions = [result["session_id"]]
        selected = retrieve(result["session_id"], "请识别节点图中的构造")
        self.assertEqual(len(selected["selected_visuals"]), 1)
        self.assertEqual(selected["input_snapshot"]["selected_visual_ids"], ["image:document:1"])
        selected_visual = selected["selected_visuals"][0]
        stored = get_visual_asset(result["session_id"], selected_visual["document_id"], selected_visual["visual_id"])
        self.assertIsNotNone(stored)
        self.assertEqual(stored.image_bytes, stream.getvalue())
        self.assertIsNone(get_visual_asset(result["session_id"], "another-document", selected_visual["visual_id"]))
        with temporary_visual_files(selected["selected_visuals"]) as paths:
            self.assertEqual(len(paths), 1)
            self.assertTrue(paths[0].exists())

    def test_up_to_four_relevant_visuals_can_be_selected_across_files(self) -> None:
        from PIL import Image

        files = []
        for index in range(4):
            stream = BytesIO()
            Image.new("RGB", (160, 90), color=(200 + index, 205, 210)).save(stream, format="PNG")
            files.append((f"节点图{index + 1}.png", stream.getvalue()))
        result = add_files(files)
        self.sessions = [result["session_id"]]
        selected = retrieve(result["session_id"], "对比这些节点图的构造")
        self.assertEqual(len(selected["selected_visuals"]), 4)
        self.assertEqual(len({item["document_id"] for item in selected["selected_visuals"]}), 4)

    def test_pdf_visual_recovery_continues_past_first_four_pages(self) -> None:
        from PIL import Image

        from backend.document_parsing.ingestion import (
            IntermediateDocument,
            SourcePointer,
            StandardFinancialDocument,
            VisualAsset,
            finalize_intermediate_evidence,
        )
        from backend.document_parsing.pdf_ingestion import PdfExtractionSummary, PdfIntakeResult

        image_stream = BytesIO()
        Image.new("RGB", (120, 80), color=(220, 220, 220)).save(image_stream, format="PNG")

        def intake(page_numbers: list[int], next_page_start: int | None) -> PdfIntakeResult:
            intermediate = finalize_intermediate_evidence(
                IntermediateDocument(
                    source_type="pdf",
                    file_name="五页扫描件.pdf",
                    parser="test-page-render",
                    tables=[],
                    visual_assets=[
                        VisualAsset(
                            visual_id=f"pdf:page:{page_number}:image",
                            kind="document_page",
                            source=SourcePointer(
                                source_type="pdf",
                                file_name="五页扫描件.pdf",
                                page_number=page_number,
                                section_id=f"pdf:page:{page_number}",
                                parser="test-page-render",
                            ),
                            media_type="image/png",
                            width_px=120,
                            height_px=80,
                            delivery_status="ready_for_vision",
                            image_bytes=image_stream.getvalue(),
                        )
                        for page_number in page_numbers
                    ],
                ),
                b"five-page-pdf-fixture",
            )
            return PdfIntakeResult(
                intermediate=intermediate,
                standard=StandardFinancialDocument(),
                validation=[],
                pdf=PdfExtractionSummary(
                    document_kind="scanned_pdf",
                    page_count=5,
                    text_characters=0,
                    page_locator="pypdf_fallback",
                    mineru_status="unavailable",
                    vision_rendered_pages=len(page_numbers),
                    vision_rendered_page_numbers=page_numbers,
                    vision_next_page_start=next_page_start,
                ),
            )

        initial = intake([1, 2, 3, 4], 5)
        final_batch = intake([5], None)
        with (
            patch.dict("os.environ", {"CUSTOMER_DOCUMENT_OCR_ENABLED": "0"}),
            patch("backend.documents.customer_sessions.ingest_uploaded_file", return_value=initial),
            patch(
                "backend.documents.customer_sessions.ingest_scanned_pdf_vision_batch",
                return_value=final_batch,
            ) as continue_batch,
        ):
            result = add_files([("五页扫描件.pdf", b"five-page-pdf-fixture")])
        continue_batch.assert_called_once_with("五页扫描件.pdf", b"five-page-pdf-fixture", page_start=5, cached_summary=initial.pdf)
        self.sessions = [result["session_id"]]
        document = result["documents"][0]
        self.assertTrue(document["visual_coverage"]["coverage_complete"])
        self.assertIsNone(document["visual_coverage"]["next_page_start"])
        self.assertEqual(document["visual_coverage"]["rendered_page_numbers"], [1, 2, 3, 4, 5])
        self.assertGreaterEqual(document["visual_coverage"]["batch_count"], 2)
        self.assertTrue(document["ready_visual_count"] >= 5)

        with (
            patch.dict("os.environ", {"CUSTOMER_DOCUMENT_OCR_ENABLED": "0"}),
            patch("backend.documents.customer_sessions.ingest_uploaded_file", return_value=initial),
            patch(
                "backend.documents.customer_sessions.ingest_scanned_pdf_vision_batch",
                side_effect=RuntimeError("render failed"),
            ),
        ):
            incomplete = add_files([("续批失败.pdf", b"five-page-pdf-fixture")])
        self.sessions.append(incomplete["session_id"])
        incomplete_coverage = incomplete["documents"][0]["visual_coverage"]
        self.assertFalse(incomplete_coverage["coverage_complete"])
        self.assertEqual(incomplete_coverage["next_page_start"], 5)
        self.assertEqual(incomplete_coverage["stop_reason"], "visual_batch_failed:RuntimeError")

    def test_citation_coverage_requires_every_selected_evidence_to_be_precise(self) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "节点参数"
        sheet.append(["部位", "安装方式"])
        sheet.append(["窗洞口", "锚固安装"])
        stream = BytesIO()
        workbook.save(stream)
        result = add_files(
            [
                ("参数.xlsx", stream.getvalue()),
                ("补充说明.txt", "窗洞口采用锚固安装。".encode("utf-8")),
            ]
        )
        self.sessions = [result["session_id"]]
        selected = retrieve(result["session_id"], "窗洞口采用什么安装方式？")
        self.assertEqual({item["document_name"] for item in selected["evidence"]}, {"参数.xlsx", "补充说明.txt"})
        self.assertGreater(selected["input_snapshot"]["precise_citation_count"], 0)
        self.assertFalse(selected["input_snapshot"]["citation_coverage_complete"])

    def test_question_window_finds_answer_after_character_6001_without_mutating_canonical(self) -> None:
        result = add_files([("超长说明.txt", "占位内容".encode("utf-8"))])
        self.sessions = [result["session_id"]]
        session = get_session(result["session_id"])
        self.assertIsNotNone(session)
        assert session is not None
        long_text = ("普通背景资料，与当前问题无关。" * 520) + "\n关键结论：窗洞口必须设置双道防水收口。"
        self.assertGreater(long_text.index("关键结论"), 6_001)
        canonical_chunk = {
            "chunk_id": "evidence:long-tail",
            "kind": "evidence",
            "text": long_text,
            "source_refs": [],
        }
        session.documents[0].chunks = [canonical_chunk]

        selected = retrieve(result["session_id"], "窗洞口的双道防水收口要求是什么？", max_text_tokens=1_200)

        self.assertTrue(any("双道防水收口" in item["text"] for item in selected["evidence"]))
        matching = next(item for item in selected["evidence"] if "双道防水收口" in item["text"])
        self.assertGreater(matching["window_offset"]["start_character"], 6_001)
        self.assertEqual(matching["original_chunk_id"], "evidence:long-tail")
        self.assertEqual(session.documents[0].chunks[0]["text"], long_text)
        snapshot = selected["input_snapshot"]
        self.assertGreater(snapshot["generated_window_count"], 1)
        self.assertEqual(snapshot["windowed_chunk_count"], 1)
        self.assertEqual(snapshot["gold_evidence_input_status"], "unknown_at_runtime")
        self.assertTrue(snapshot["canonical_evidence_preserved"])

    def test_language_aware_token_estimate_handles_chinese_and_english(self) -> None:
        self.assertEqual(_estimate_text_tokens("中文测试"), 4)
        self.assertEqual(_estimate_text_tokens("abcdefgh"), 2)
        self.assertEqual(_estimate_text_tokens("hello world"), 4)
        bilingual = _estimate_text_tokens("外墙 insulation system 2026")
        self.assertGreaterEqual(bilingual, 7)
        self.assertLess(bilingual, len("外墙 insulation system 2026"))

    def test_metadata_excluded_and_ocr_candidate_keeps_page_and_unverified_role(self) -> None:
        from backend.documents.customer_sessions import StoredVisualAsset
        result = add_files([('candidate.txt', b'placeholder')])
        self.sessions = [result['session_id']]
        document = get_session(result['session_id']).documents[0]
        metadata = '[BLOCK page=9 type=no_text parser=pdf]\n[VISUAL id=page9]'
        document.chunks = [dict(chunk_id='empty', kind='evidence', text=metadata, source_refs=[])]
        document.source_locations['page9'] = dict(page_number=9)
        document.visuals = [StoredVisualAsset(visual_id='page9', document_id=document.document_id,
            document_name=document.file_name, kind='pdf_page', source=dict(page_number=9),
            media_type='image/png', metadata=dict(ocr_literal_text='Fieldwork personnel: Wave I 52, Wave II 50'),
            image_bytes=None, searchable_text='Fieldwork personnel: Wave I 52, Wave II 50')]
        found = retrieve(result['session_id'], 'fieldwork personnel', enable_semantic_rerank=False)
        self.assertTrue(found['evidence'])
        candidate = found['evidence'][0]
        self.assertIn('Wave I 52', candidate['text'])
        self.assertFalse(candidate['facts_eligible'])
        self.assertTrue(candidate['candidate_requires_pixel_verification'])
        self.assertEqual(candidate['citations'][0]['source_page'], 9)
        self.assertEqual(found['input_snapshot']['metadata_only_chunks_excluded_from_answers'], 1)
        self.assertEqual(document.chunks[0]['text'], metadata)

    def test_repeated_retrieval_reuses_windows_but_not_mutable_rank_state(self) -> None:
        result = add_files([('cache.txt', b'BeamY thickness 21 mm. Installation uses anchors.')])
        self.sessions = [result['session_id']]
        retrieve(result['session_id'], 'thickness', enable_semantic_rerank=False)
        with patch('backend.documents.customer_sessions._window_chunk', side_effect=AssertionError('unexpected resplit')):
            found = retrieve(result['session_id'], 'anchors', enable_semantic_rerank=False)
        self.assertTrue(any('anchors' in e['text'] for e in found['evidence']))
        document = get_session(result['session_id']).documents[0]
        self.assertLessEqual(document.window_cache_bytes, 1024**2)
        self.assertTrue(all('score' not in window for windows in document.window_cache.values() for window in windows))

    def test_interleaved_visual_geometry_does_not_displace_prose_or_break_offsets(self) -> None:
        from backend.documents.customer_sessions import _window_chunk
        text = 'Purpose: flexible working application.\n[VISUAL id=logo metadata=' + ('coordinate ' * 300) + ']\nApproval requires an authorised manager.'
        windows = _window_chunk(text, 'C', window_tokens=64, overlap_tokens=8)
        self.assertTrue(any('flexible working' in w['text'] for w in windows))
        self.assertTrue(any('authorised manager' in w['text'] for w in windows))
        self.assertFalse(any('coordinate' in w['text'] for w in windows))
        for window in windows:
            self.assertEqual(text[window['start_character']:window['end_character']].strip(), window['text'].strip())

    def test_four_file_budget_is_reserved_only_for_relevant_documents(self) -> None:
        relevant_a = ("窗洞口防水节点应连续密封。" * 180).encode("utf-8")
        relevant_b = ("窗洞口锚固节点采用机械固定。" * 180).encode("utf-8")
        result = add_files(
            [
                ("防水说明.txt", relevant_a),
                ("锚固说明.txt", relevant_b),
                ("会议纪要.txt", ("参会人员与会议时间。" * 180).encode("utf-8")),
                ("库存清单.txt", ("办公用品库存盘点。" * 180).encode("utf-8")),
            ]
        )
        self.sessions = [result["session_id"]]
        with patch.dict(
            "os.environ",
            {
                "CUSTOMER_ATTACHMENT_TEXT_TOKEN_BUDGET": "1000",
                "CUSTOMER_ATTACHMENT_TEXT_TOKEN_SAFETY_MAX": "1200",
                "CUSTOMER_ATTACHMENT_WINDOW_TOKENS": "256",
                "CUSTOMER_ATTACHMENT_WINDOW_OVERLAP_TOKENS": "32",
                "CUSTOMER_ATTACHMENT_DOCUMENT_MIN_TOKENS": "200",
            },
        ):
            selected = retrieve(result["session_id"], "窗洞口防水和锚固节点如何处理？")

        snapshot = selected["input_snapshot"]
        self.assertLessEqual(snapshot["selected_estimated_text_tokens"], 1_000)
        self.assertEqual(
            {item["document_name"] for item in selected["evidence"]},
            {"防水说明.txt", "锚固说明.txt"},
        )
        quota_by_name = {item["document_name"]: item for item in snapshot["per_document_quotas"]}
        self.assertGreaterEqual(quota_by_name["防水说明.txt"]["selected_estimated_tokens"], 200)
        self.assertGreaterEqual(quota_by_name["锚固说明.txt"]["selected_estimated_tokens"], 200)
        self.assertEqual(quota_by_name["会议纪要.txt"]["selected_estimated_tokens"], 0)
        self.assertEqual(quota_by_name["库存清单.txt"]["selected_estimated_tokens"], 0)
        self.assertTrue(quota_by_name["会议纪要.txt"]["zero_relevance"])
        self.assertEqual(snapshot["selection_policy"], "relevant_documents_minimum_then_global_window_ranking")


if __name__ == "__main__":
    unittest.main()
